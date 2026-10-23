"""Render ảnh và kiểm tra riêng Bộ-câu-hỏi, không sửa tài liệu cũ."""
from pathlib import Path
import json,re,hashlib,argparse
from datetime import date
from docx import Document
import pypdfium2 as pdfium

parser=argparse.ArgumentParser();parser.add_argument('--render',action='store_true');parser.add_argument('--visual-reviewed',action='store_true');args=parser.parse_args()
root=Path(__file__).resolve().parents[1];out=root/'tmp/defense/render'
docpath=root/'Bộ-câu-hỏi.docx';doc=Document(docpath)
index=json.loads((root/'docs/defense-question-index.json').read_text(encoding='utf-8'))
preserved=json.loads((root/'docs/defense-preservation.json').read_text(encoding='utf-8'))
for file,h in preserved.items():assert hashlib.sha256((root/file).read_bytes()).hexdigest()==h, 'Protected file changed: '+file
ids=[f'A{i:02}' for i in range(1,21)]+[f'B{i:02}' for i in range(1,31)]+[f'C{i:02}' for i in range(1,11)]
headings=[p.text for p in doc.paragraphs if p.style.name=='Heading 2']
assert [h[:3] for h in headings]==ids
assert len(doc.inline_shapes)==3
assert {p.name for p in root.glob('*.docx')}=={'BaoCao_Assignment2_PC-Astral-Market.docx','UseCase_Logic_Ham_PC-Astral-Market.docx','Bộ-câu-hỏi.docx'}
assert {r['question'] for r in index['references']}==set(ids)
for ref in index['references']:
 p=root/ref['file'] if ref['file'].startswith('tools/') else root/'PC-Astral-Market'/ref['file']
 lines=p.read_text(encoding='utf-8').splitlines()
 assert ref['anchor'] in lines[ref['line']-1],ref
 if ref['function']:
  f=ref['function'];assert f['start']<=ref['line']<=f['end'];assert f['name'].split('.')[-1] in lines[f['start']-1]
for snippet in index['snippets']:
 lines=(root/'PC-Astral-Market'/snippet['file']).read_text(encoding='utf-8').splitlines()
 assert lines[snippet['start']-1:snippet['start']-1+snippet['count']]==snippet['source']
 assert 3<=snippet['count']<=10
pdf=pdfium.PdfDocument(out/'questions.pdf');texts=[];toc={};body={};bounds=[]
if args.render:
 for file in out.glob('page-*.png'):file.unlink()
for i,page in enumerate(pdf):
 textpage=page.get_textpage();text=textpage.get_text_range();texts.append(text)
 if args.render:page.render(scale=1.7).to_pil().save(out/f'page-{i+1}.png')
 for line in text.splitlines():
  match=re.match(r'^([ABC]\d{2})\s+.*?\.{3,}\s*(\d+)\s*$',line.strip())
  if match:toc[match[1]]=int(match[2])
  elif re.match(r'^[ABC]\d{2}\s+',line.strip()):
   identifier=line.strip()[:3]
   if identifier not in body and i>3:body[identifier]=i+1
 assert len(text.strip())>250 or i==5, f'Almost blank page {i+1}'
 for n in range(textpage.count_chars()):
  ch=textpage.get_text_range(n,1)
  if not ch.strip():continue
  left,bottom,right,top=textpage.get_charbox(n)
  if left<75 or right>550 or bottom<20 or top>825:bounds.append([i+1,n,left,bottom,right,top])
assert len(toc)==60,(len(toc),set(ids)-set(toc))
assert set(body)==set(ids),(set(ids)-set(body))
assert toc==body,[(k,toc[k],body.get(k)) for k in ids if toc[k]!=body.get(k)]
assert not bounds, bounds[:10]
tests=json.loads((root/'docs/defense-practice-tests.json').read_text(encoding='utf-8'))
assert len(tests['results'])==11 and all(x['status']=='PASS' for x in tests['results'])
summary={'date':date.today().isoformat(),'file':docpath.name,'pages':len(pdf),'questions':60,'core':20,'advanced':30,'practice':10,'diagrams':3,
 'source_references_checked':len(index['references']),'source_excerpts_checked':len(index['snippets']),
 'toc_question_pages_checked':len(toc),'practice_tests_passed':11,'protected_files_unchanged':len(preserved),
 'sha256':hashlib.sha256(docpath.read_bytes()).hexdigest(),'renderer':'Microsoft Word COM PDF export + pypdfium2; bundled renderer unavailable because LibreOffice is absent.',
 'all_pages_visually_reviewed':args.visual_reviewed,'status':'PASS' if args.visual_reviewed else 'STRUCTURAL_PASS_PENDING_VISUAL'}
(root/'docs/defense-document-qa.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(summary,ensure_ascii=True,indent=2))

from pathlib import Path
import json,hashlib,sys
from docx import Document
import pypdfium2 as pdfium
root=Path(__file__).resolve().parents[1]
index=json.loads((root/'docs/js-revision-index.json').read_text(encoding='utf-8'))
for ref in index['references']:
 lines=(root/'PC-Astral-Market'/ref['file']).read_text(encoding='utf-8').splitlines()
 assert ref['anchor'] in lines[ref['line']-1]
protected=json.loads((root/'docs/defense-preservation.json').read_text(encoding='utf-8'))
for file,sha in protected.items():assert hashlib.sha256((root/file).read_bytes()).hexdigest()==sha,file
doc=Document(root/'Bộ-câu-hỏi.docx')
assert len([p for p in doc.paragraphs if p.style.name=='Heading 2'])==27
pdf=pdfium.PdfDocument(root/'tmp/defense/render/questions.pdf')
texts=[]
for i,page in enumerate(pdf):
 page.render(scale=1.5).to_pil().save(root/f'tmp/defense/render/js-{i+1}.png')
 text=page.get_textpage();texts.append(text.get_text_range())
 for j in range(text.count_chars()):
  if not text.get_text_range(j,1).strip():continue
  left,bottom,right,top=text.get_charbox(j)
  assert left>=75 and right<=550 and bottom>=20 and top<=825,(i,left,bottom,right,top)
assert len(pdf)==10
reviewed='--visual-reviewed' in sys.argv
report={'pages':len(pdf),'questions':27,'topics':9,'references_checked':len(index['references']),'protected_files_unchanged':len(protected),'visually_reviewed':reviewed,'status':'PASS' if reviewed else 'PENDING_VISUAL','sha256':hashlib.sha256((root/'Bộ-câu-hỏi.docx').read_bytes()).hexdigest()}
(root/'docs/js-revision-qa.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report))

"""Tạo riêng Bộ-câu-hỏi.docx; không sửa website và hai tài liệu hiện có."""
from pathlib import Path
from datetime import date
import json, hashlib, textwrap
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from defense_content import QUESTIONS, M,U,S,E,D,H,P,G,O,A,r
from defense_exercises import EXERCISES

ROOT=Path(__file__).resolve().parents[1];SITE=ROOT/'PC-Astral-Market'
FIG=ROOT/'docs/diagrams/defense';FIG.mkdir(parents=True,exist_ok=True)
OUT=ROOT/'Bộ-câu-hỏi.docx';refs_used=[];snippets=[]
SYMBOLS=json.loads((ROOT/'docs/defense-source-symbols.json').read_text(encoding='utf-8'))

def source(file):return ROOT/file if file.startswith('tools/') else SITE/file
def locate(file,needle):
    lines=source(file).read_text(encoding='utf-8').splitlines()
    matches=[i+1 for i,line in enumerate(lines) if needle in line]
    assert matches,(file,needle)
    return matches[0],lines
def references(refs,question):
    groups={}
    for file,anchor,label in refs:
        line,_=locate(file,anchor)
        candidates=[x for x in SYMBOLS.get(file,[]) if x['start']<=line<=x['end']]
        function=min(candidates,key=lambda x:x['end']-x['start']) if candidates and not anchor.strip().startswith('class ') else None
        location=f'{label}, dòng {line}'
        if function:
            location=f"{function['name']}, dòng {function['start']}"
            if function['start']!=line:location+=f' → {line}'
        if file.startswith('tools/') and anchor.startswith('def '):location=anchor[4:].split('(')[0]+f' (dòng {line})'
        refs_used.append(dict(question=question,file=file,anchor=anchor,line=line,function=function))
        groups.setdefault(file,[]).append(location)
    return ' | '.join(file+': '+'; '.join(values) for file,values in groups.items())+'.'

def configure():
    d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.left_margin=Cm(3);s.right_margin=s.top_margin=s.bottom_margin=Cm(2)
    s.header_distance=s.footer_distance=Cm(.8)
    for name,size in [('Normal',12),('Title',21),('Subtitle',12),('Heading 1',15),('Heading 2',12),('Caption',10),('TOC 1',11),('TOC 2',10.5),('Code',10)]:
        if name not in d.styles:d.styles.add_style(name,1)
        st=d.styles[name];st.font.name='Consolas' if name=='Code' else 'Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        pr=st.element.get_or_add_rPr();fonts=pr.find(qn('w:rFonts'))
        if fonts is None:fonts=OxmlElement('w:rFonts');pr.append(fonts)
        for a in list(fonts.attrib):
            if 'theme' in a.lower():del fonts.attrib[a]
        for a in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+a),st.font.name)
        st.paragraph_format.line_spacing=1.15;st.paragraph_format.space_after=Pt(4)
        if name.startswith('Heading'):st.font.bold=True;st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(8)
        if name.startswith('TOC'):
            # Word dùng tên built-in viết thường; tránh tạo TOC 11/21 rồi bỏ qua kiểu.
            st.element.attrib.pop(qn('w:customStyle'),None)
            st.element.find(qn('w:name')).set(qn('w:val'),name.lower())
            st.paragraph_format.line_spacing=1;st.paragraph_format.space_after=Pt(1)
            st.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
            if name=='TOC 2':st.paragraph_format.left_indent=Cm(.6)
        if name=='Code':st.paragraph_format.line_spacing=1;st.paragraph_format.space_after=Pt(4)
    for st in d.styles:
        for border in st.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    s.header.paragraphs[0].text='PC ASTRAL MARKET  |  ÔN BẢO VỆ WEB1044'
    s.header.paragraphs[0].runs[0].font.size=Pt(9)
    f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    f.add_run('Tài liệu luyện tập  •  ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
    return d

doc=configure()
def para(text='',label=None,style=None):
    p=doc.add_paragraph(style=style)
    if label:p.add_run(label+': ').bold=True
    p.add_run(text);return p
def newpage(title=None):
    doc.add_page_break()
    if title:doc.add_heading(title,1)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        edge=OxmlElement('w:'+side)
        for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:edge.set(qn('w:'+key),value)
        borders.append(edge)
    t._tbl.tblPr.append(borders)
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    for row in rows:
        cells=t.add_row().cells
        for i,value in enumerate(row):cells[i].text=str(value)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for i,c in enumerate(row.cells):
            c.width=Cm(widths[i]);tc=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF2' if ri==0 else 'FFFFFF');tc.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1
                for run in p.runs:run.font.size=Pt(10.5);run.bold=ri==0
    return t
def code(text):
    p=para(text,style='Code');p.paragraph_format.keep_together=True
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F4F6');p._p.get_or_add_pPr().append(sh)
    return p
def diagram(name,rows):
    width=1400;height=110+len(rows)*190
    im=Image.new('RGB',(width,height),'white');dr=ImageDraw.Draw(im);font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',32)
    for ri,labels in enumerate(rows):
        step=(width-40)/len(labels);y=30+ri*190
        for i,label in enumerate(labels):
            x=20+i*step;right=x+step-32
            dr.rounded_rectangle((x,y,right,y+140),radius=9,fill='#F2F5F7',outline='#556574',width=2)
            for j,line in enumerate(label.split('\n')):
                assert dr.textlength(line,font=font)<right-x-12,(name,line)
                dr.text(((x+right)/2,y+35+j*37),line,font=font,anchor='mm',fill='black')
            if i<len(labels)-1:
                dr.line((right,y+70,right+26,y+70),fill='#556574',width=3)
                dr.polygon([(right+26,y+70),(right+17,y+64),(right+17,y+76)],fill='#556574')
    p=FIG/(name+'.png');im.save(p);return p
def figure(path,caption):
    p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(path),width=Cm(16))
    para(caption,style='Caption')

para('BỘ CÂU HỎI BẢO VỆ\nPC ASTRAL MARKET',style='Title')
para('Môn WEB1044  •  60 câu có đáp án và vị trí mã',style='Subtitle')
para('Tài liệu giúp sinh viên luyện trả lời, mở đúng mã và thực hành sửa nhỏ trước buổi bảo vệ. Câu hỏi được viết theo hai vai giảng viên và sinh viên, dựa trên phiên bản dự án hiện tại. Đây không phải đề thi hoặc rubric chính thức của FPT Polytechnic.')
para('20 câu trọng tâm  •  30 câu đào sâu  •  10 bài sửa tại chỗ',label='Lộ trình')
para('Ngày đối chiếu mã: '+date.today().strftime('%d/%m/%Y')+'. Tên file và số dòng tính từ thư mục PC-Astral-Market, riêng tools/ tính từ thư mục dự án. Khi mã đổi, tìm theo tên hàm trước rồi kiểm tra lại số dòng.')
para('Nếu ít thời gian, ôn A01–A20 trước. Mỗi câu: đọc câu hỏi → che đáp án và nói thành tiếng → mở mã → thực hiện phần chứng minh. Chỉ đánh dấu tự làm được khi vừa nói rõ vừa thao tác đúng.')
para('20 câu cần ôn trước',label='Tra nhanh')
table(['A01 đến A10','A11 đến A20'],[[f"{QUESTIONS[i]['id']}  {QUESTIONS[i]['title']}",f"{QUESTIONS[i+10]['id']}  {QUESTIONS[i+10]['title']}"] for i in range(10)],[8,8])
para('Giữ bản nộp nguyên trạng. Các bài sửa C01–C10 chỉ thực hiện trên bản sao; không dùng thông tin tài khoản hoặc mật khẩu cá nhân.',label='Cách luyện')

newpage();para('MỤC LỤC',style='Title')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u');doc.add_paragraph()._p.append(field)

newpage('1 Cách demo và lần theo mã')
para('Ưu tiên mở bằng Live Server đúng thư mục website. Giữ nguyên địa chỉ, cổng và hồ sơ trình duyệt trong buổi luyện để tránh nhầm vùng lưu LocalStorage. Website cũng có phương án mở file trực tiếp, nhưng không coi mọi trình duyệt xử lý file:// giống nhau.')
table(['Thời gian','Thao tác và lời trình bày'],[
('0:00–0:30','Giới thiệu cửa hàng máy tính mô phỏng, bốn danh mục và ba nhóm sản phẩm.'),
('0:30–1:15','Tìm ASUS, lọc Laptop, sắp xếp giá; giải thích ba chức năng tự đề xuất.'),
('1:15–2:00','Xem chi tiết, thêm hai sản phẩm; thêm trùng một món để chứng minh gộp dòng.'),
('2:00–2:45','Đổi số lượng, xóa một dòng, F5; đối chiếu số lượng và tổng tiền.'),
('2:45–3:30','Mở form; gửi trống để thấy lỗi, điền dữ liệu thử. Nếu giá thay đổi, đọc cảnh báo và xác nhận lại.'),
('3:30–4:30','Đặt đơn, chỉ mã và tổng, F5 xác nhận. Nói rõ không thu tiền hoặc gửi cửa hàng thật.'),
('Dự phòng','Để khoảng 30 giây xử lý tải trang hoặc giải thích câu hỏi, không bấm liên tiếp theo quán tính.')],[2.5,13.5])
para('Demo tài khoản bổ sung khoảng 1–2 phút: đăng ký onthi01 bằng mật khẩu thử Demo2026! → đăng nhập → F5 → đăng xuất → đăng nhập lại. Chỉ xem username/salt/passwordHash trong storage; kết thúc demo chỉ bấm khi chấp nhận xóa dữ liệu thử.',label='Tài khoản')
para('Đọc trước ba điều dễ nhầm: F5 không tự xóa LocalStorage; tổng tiền hiện dùng for chứ không dùng reduce; đăng nhập này không phải xác thực có backend.',label='Cần nói đúng')
para('Nếu demo lỗi: dừng thao tác, đọc thông báo, mở Console/Network và kiểm tra URL. Không sửa nhiều file hoặc xóa toàn bộ storage chỉ để đoán lỗi. Dùng bản sao dự phòng và mô tả trung thực điều quan sát được.',label='Khi gặp lỗi')

newpage('2 Bản đồ mã và cách chứng minh')
table(['Chức năng','Nơi mở và điều cần chỉ'],[
('Dữ liệu','js/products.json: 24 sản phẩm; js/du-lieu.js: danh mục, banner, thời hạn và nạp JSON.'),
('Mô hình','js/mo-hinh.js: Product, Cart; giá, số lượng và tổng tiền.'),
('Lưu và tài khoản','js/luu-tru.js: khóa lưu, JSON, đăng ký, băm mật khẩu, đăng nhập và reset.'),
('Hiển thị và lọc','js/hien-thi.js: thẻ sản phẩm, lọc, so sánh, validation và giao diện chung.'),
('Khởi tạo','js/su-kien.js: tạo Product/Cart, listener chung và đồng bộ tab.'),
('Logic trang','js/trang/: trang-chu, san-pham, chi-tiet, gio-hang, dat-hang, tai-khoan.'),
('Bố cục','HTML ở gốc; css/style.css; ảnh img/; font local trong assets/fonts/.'),
('Mở trực tiếp','js/san-pham-du-phong.js tự sinh từ JSON; công cụ tools/sync_offline_products.py đồng bộ.')],[3.3,12.7])
para('Trong VS Code, mở file bằng Ctrl+P, nhập đường dẫn; Ctrl+G để đến số dòng. Nếu dòng đã dịch chuyển, Ctrl+F tên hàm. Đọc cả bên gọi và giá trị trả về, không chỉ thuộc tên hàm.',label='Mở đúng mã')
para('F12 → Application → Local Storage → chọn đúng origin. Đọc các khóa pc-astral.*; Console xem lỗi JavaScript, Network xem tài nguyên thiếu. Dùng phiên thử riêng nếu sửa/xóa khóa để luyện lỗi; không dùng tài khoản thật.',label='Xem dữ liệu')
para('Với 5–10 phút giải thích: dành 1 phút nêu đầu vào, 2 phút chỉ nơi kiểm tra, 2 phút lần theo Cart/giá/lưu trữ, 1–2 phút chỉ render và phần còn lại cho ca lỗi. Ưu tiên các câu A12–A20 và B26–B29.',label='Giải thích mã')
para('“Chỉ vào mã” ghi tên hàm và dòng khai báo. Ví dụ dòng 204 → 224 nghĩa là hàm khai báo ở 204, câu lệnh cần chỉ ở 224. “Mẫu sửa” chỉ dành cho bản luyện. Dòng trích có số thật, không chép số dòng vào JS; phải mở toàn hàm để hiểu ngữ cảnh.',label='Đọc trích đoạn')

newpage('3 Ba luồng cần nhớ')
figure(diagram('tai-khoan',[
['Đăng ký\nKiểm tra form','Tên chưa trùng\nTạo salt và hash','Kiểm tra reset\nLưu tài khoản','Thông báo\nMời đăng nhập'],
['Đăng nhập\nTìm username','Băm mật khẩu\nVới salt đã lưu','Kiểm tra lại\nSo passwordHash','Lưu AUTH_KEY\nCập nhật tên']]),'Hình 1. Tài khoản chỉ lưu username, salt và hash; lỗi ở bước kiểm tra hoặc lưu thì không báo thành công.')
figure(diagram('gio-hang',[
['Nhận thao tác\nKiểm tra số lượng','Cập nhật Cart\nGộp theo ID','Lưu JSON\nTính tổng','Cập nhật dòng\nVà badge'],
['F5 trang\nĐọc LocalStorage','Kiểm tra dòng\nBỏ dữ liệu sai','Tạo lại Cart\nTừ dữ liệu lưu','Hiển thị giỏ\nTheo Cart mới']]),'Hình 2. Nếu ghi storage lỗi, bộ nhớ và dữ liệu lưu có thể khác; đọc cảnh báo, không hứa F5 sẽ giữ thao tác.')
figure(diagram('dat-hang',[
['Kiểm tra form\nGiỏ và giá','Lưu bản chụp đơn\nTrước khi dọn giỏ','Xóa bản nháp\nThử dọn giỏ','Hiển thị mã đơn\nĐã được lưu']]),'Hình 3. Giá/giỏ đổi: xin xác nhận lại. Lưu đơn lỗi: giữ form/giỏ. Dọn giỏ lỗi: chỉ thử dọn đúng đơn đã lưu.')

def add_question(item):
    start=len(doc.paragraphs)
    doc.add_heading(item['id']+' '+item['title'],2)
    para(item['ask'],label='Giảng viên hỏi')
    para(item['answer'],label='Em trả lời')
    para(item['logic'],label='Logic')
    para(references(item['refs'],item['id']),label='Chỉ vào mã')
    if item.get('excerpt'):
        file,anchor,count=item['excerpt'][:3];line,lines=locate(file,anchor)
        if len(item['excerpt'])==4:line+=item['excerpt'][3]
        raw=lines[line-1:line-1+count];common=min((len(x)-len(x.lstrip()) for x in raw if x.strip()),default=0)
        text='\n'.join(f'{line+i:>3}  {x[common:]}' for i,x in enumerate(raw))
        code(text);snippets.append(dict(question=item['id'],file=file,start=line,count=count,source=raw))
    para(item['demo'],label='Chứng minh')
    para(item['follow'],label='Hỏi tiếp')
    para(item['trap'],label='Tránh nhầm')
    # Giữ mỗi câu thành một khối nếu còn vừa một trang; tránh tiêu đề lẻ.
    for p in doc.paragraphs[start:-1]:p.paragraph_format.keep_with_next=True

newpage('4 Hai mươi câu trọng tâm')
for i,item in enumerate(QUESTIONS[:20]):
    if i and i%2==0:doc.add_page_break()
    add_question(item)
newpage('5 Ba mươi câu đào sâu')
for i,item in enumerate(QUESTIONS[20:]):
    if i and i%2==0:doc.add_page_break()
    add_question(item)

newpage('6 Mười bài sửa tại chỗ')
para('Các mẫu sau là đáp án luyện tập, chưa áp dụng vào website chính. Tạo bản sao, sửa đúng yêu cầu, chạy thử và hoàn tác. Trước khi thi, kiểm tra có sẵn môi trường Python nếu cần sinh JSON dự phòng; nếu chưa có, luyện bằng HTTP trước và chuẩn bị bản đóng gói đã đồng bộ.')
for i,item in enumerate(EXERCISES):
    doc.add_heading(item['id']+' '+item['title'],2)
    para(item['ask'],label='Giảng viên giao bài')
    para(item['answer'],label='Em trình bày cách làm')
    para(references(item['refs'],item['id']),label='Chỉ vào mã gốc')
    para(item['steps'],label='Các bước sửa')
    if item.get('solution'):
        para('Đoạn bổ sung hoặc thay thế trong bản luyện; không phải toàn bộ hàm.',label='Mẫu sửa').paragraph_format.keep_with_next=True
        code(item['solution'])
    para(item['checks'],label='Kiểm tra và kết quả mong đợi')
    para(item['follow'],label='Hỏi tiếp')
    para(item['trap'],label='Dễ mất điểm')
    para(item['undo'],label='Hoàn tác')
    para('Đã chạy bài tương ứng trên bản sao bằng Chrome; xem kết quả tổng hợp cuối tài liệu. Đây không phải thay đổi của website nộp bài.',label='Minh chứng luyện')

newpage('7 Luyện hội đồng và tự đánh giá')
table(['Buổi luyện','Câu và trình tự'],[
('Đề 1  Giỏ hàng','Demo mua hàng 3–5 phút; giải thích A08, A12, A13, A15; tình huống B17, B20; sửa C03.'),
('Đề 2  Tài khoản','Demo đăng ký/đăng nhập/F5; giải thích A16–A19; tình huống B21, B24, B25; sửa C09.'),
('Đề 3  Dữ liệu và giá','Demo tìm/lọc/sắp xếp và form; giải thích A05, B08–B10; tình huống B14, B27–B30; sửa C05 hoặc C08 theo người hỏi.')],[3.6,12.4])
para('Mỗi lượt dành khoảng 3–5 phút demo, 5–10 phút giải thích và 3–5 phút bài sửa. Nhóm ba người luân phiên làm sinh viên, giảng viên hỏi và người ghi nhận. Đây là cách luyện đề xuất, không phải thời lượng thi chính thức.',label='Tổ chức nhóm')
table(['Tiêu chí tự đánh giá','Chưa nắm','Có gợi ý','Tự làm được'],[
('Chạy đủ luồng mua hàng và nói đúng giới hạn','□','□','□'),
('Giải thích đầu vào, điều kiện, xử lý, kết quả','□','□','□'),
('Mở đúng file/hàm và chỉ được dòng xử lý','□','□','□'),
('Giải thích dữ liệu lưu, F5, mật khẩu và reset','□','□','□'),
('Sửa đúng yêu cầu, thử biên và hoàn tác','□','□','□'),
('Tự chẩn đoán bằng Console/Network/Storage','□','□','□')],[9.1,2.2,2.2,2.5])
para('Chỉ đánh dấu sau khi tự thực hiện; không suy ra khả năng bảo vệ của sinh viên từ việc bộ test chạy thành công. Cả ba thành viên cần nắm luồng chính; không tự nhận phần việc đã làm thay cho bạn khác.',label='Cách dùng bảng')
table(['Tiêu chí trong ảnh người dùng cung cấp','Tỷ trọng gợi ý'],[
('Chức năng chạy đúng, đủ luồng mua hàng','15%'),('Giải thích được mã của mình','15%'),('Tổ chức mã rõ ràng, tách file','5%'),('Giao diện gọn gàng, dùng được','5%')],[12,4])
para('Bốn tỷ trọng cộng thành 40%, không phải thang 100% hoàn chỉnh. Đây chỉ là bảng gợi ý trong ảnh; giảng viên quyết định rubric chính thức. Không gắn các tỷ trọng này thành điểm đã được hội đồng xác nhận.',label='Đọc đúng bảng điểm')

newpage('8 Tránh mất điểm và chuẩn bị trước buổi thi')
table(['Lỗi','Dấu hiệu và cách xử lý'],[
('Chức năng chạy nửa vời','Có nút nhưng thao tác cuối không hoàn thành. Chạy trọn một luồng từ sản phẩm đến xác nhận trước khi thêm chức năng phụ; nói rõ phần đang mô phỏng.'),
('Giỏ lệch số liệu','Badge, dòng giỏ, tổng hoặc F5 không khớp. Kiểm tra Cart trước, rồi dữ liệu lưu, cuối cùng DOM; sửa dữ liệu rồi hiển thị, không chỉ sửa chữ tổng.'),
('Console đầy lỗi đỏ','Có exception hoặc tài nguyên thiếu. Mở F12, làm lại thao tác gây lỗi, đọc file/dòng và sửa nguyên nhân; không chỉ xóa log hoặc che thông báo.')],[4,12])
for text in [
'□ Chạy toàn bộ thư mục trên máy khác hoặc đường dẫn có dấu/khoảng trắng; không copy riêng HTML.',
'□ Chọn một địa chỉ HTTP cố định, kiểm tra ảnh/font local và cổng đang phục vụ đúng website.',
'□ Kiểm tra ngày hết khuyến mãi trước buổi thi. Nếu đã hết, giải thích trạng thái hoặc luyện thời hạn trên bản sao; không giả vờ đang giảm giá.',
'□ Ôn A01–A20 thành tiếng và mở được mã trong khoảng 10–20 giây; số dòng có thể đổi nếu sửa file.',
'□ Chuẩn bị dữ liệu tài khoản thử và giỏ demo; không dùng mật khẩu cá nhân.',
'□ Tạo bản sao để luyện C01–C10; biết cách hoàn tác và cách đồng bộ JSON dự phòng.',
'□ Mở sẵn Console và các file Product/Cart, storage, đặt hàng; đóng tab/dự án không liên quan.',
'□ Thử form sai, form đúng, F5 và reset; không bấm Kết thúc demo trước khi trình bày dữ liệu cần giữ.',
'□ Cả ba thành viên luân phiên trả lời và demo; xác định phần chưa hiểu để ôn lại.'
]:para(text)
para('Nếu chưa biết câu trả lời: nói rõ phần em đã xác nhận, mở hàm liên quan và lần theo dữ liệu. Đừng đoán rằng hàm dùng reduce, có backend hoặc F5 xóa tài khoản khi mã không như vậy.',label='Khi bị hỏi bất ngờ')

newpage('9 Kết quả đối chiếu và nguồn tài liệu')
tests=json.loads((ROOT/'docs/defense-practice-tests.json').read_text(encoding='utf-8'))
assert len(tests['results'])==11 and all(x['status']=='PASS' for x in tests['results']), 'Practice tests must pass before authoring final claims'
para('Ngày kiểm tra '+date.today().strftime('%d/%m/%Y')+'. Các bài C01–C10 được thực hiện trên mười bản sao riêng trong đường dẫn có dấu và khoảng trắng. Website nguồn, báo cáo Y3, tài liệu Use Case và ZIP nộp bài được bảo vệ bằng đối chiếu SHA-256 trước/sau.')
table(['Phần đã chạy','Kết quả và giới hạn'],[
('C01–C03','Thêm sản phẩm và bản file://; giá trên danh sách/chi tiết/giỏ/đơn; giới hạn kho và số lượng sai: đạt.'),
('C04–C05','Banner thứ tư, bộ đếm và vòng chuyển; hết hạn giữa lúc nhập đơn, giữ form và xác nhận lại: đạt.'),
('C06–C08','Tìm theo mã; cận giá; tên Z–A và trạng thái sau F5, không đổi mảng gốc: đạt.'),
('C09–C10','Nháp 301 ký tự bị chặn, sửa 300 đặt được; tái hiện 404 ảnh rồi sửa đường dẫn và tải đúng: đạt.'),
('Demo bổ sung','Mở đủ bảy trang; đăng ký, đăng nhập, F5, đăng xuất/đăng nhập lại, xem các trường lưu và reset: đạt.'),
('Môi trường','Google Chrome qua HTTP local; C01 kiểm tra thêm file://. C05 dùng đồng hồ kiểm thử. Không coi đây là kiểm tra mọi trình duyệt hoặc mọi tình huống.')],[3.2,12.8])
para('Nguồn chính: mã PC-Astral Market hiện tại; báo cáo Assignment 2 và tài liệu Use Case, Logic và Hàm trong dự án. Bốn ảnh do người dùng cung cấp được dùng để xây dựng cấu trúc demo, câu hỏi thường gặp, tự đánh giá và ba lỗi tránh mất điểm.',label='Nguồn nội dung')
para('Kho mã tham khảo của nhóm: https://github.com/ngminhtamtech-cmd/Assignment-WEB1044. Số dòng trong tài liệu lấy từ bản local tại ngày đối chiếu, không mặc định rằng kho GitHub đã chứa mọi thay đổi local.',label='GitHub')
para('Dữ liệu câu hỏi: tools/defense_content.py; bài luyện: tools/defense_exercises.py; dựng Word: tools/build_defense_document.py; kiểm thử: tools/test_defense.cjs. Danh sách tham chiếu và kết quả QA được giữ trong docs/ để có thể kiểm tra lại khi mã thay đổi.',label='Nguồn tái tạo')
para('Không cần học thuộc toàn bộ 60 đáp án. Mục tiêu là tự diễn đạt đúng logic, mở được mã và chứng minh bằng thao tác; phần nào chưa làm được thì quay lại đọc bên gọi, dữ liệu đầu vào và kết quả trả về.',label='Mục tiêu ôn tập')

assert len(QUESTIONS)==50 and len(EXERCISES)==10
doc.save(OUT)
manifest={'date':date.today().isoformat(),'document':OUT.name,'questions':[{'id':x['id'],'title':x['title']} for x in QUESTIONS+EXERCISES],
 'references':refs_used,'snippets':snippets,'source_sha256':{p.relative_to(ROOT).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() for p in SITE.rglob('*') if p.is_file()},'figures':3}
(ROOT/'docs/defense-question-index.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print('Created 60 questions with',len(refs_used),'verified source anchors.')

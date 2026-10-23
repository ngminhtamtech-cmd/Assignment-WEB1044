"""Kiểm tra Word; chỉ xác nhận QA sau khi xem mọi PNG."""
from pathlib import Path
import json, hashlib, re, argparse
from datetime import date
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
import pypdfium2 as pdfium

parser=argparse.ArgumentParser()
parser.add_argument('--visual-reviewed',action='store_true',help='Chỉ dùng sau khi xem tất cả trang')
args=parser.parse_args()
root=Path(__file__).resolve().parents[1]
state=json.loads((root/'docs/report-content-sync.json').read_text(encoding='utf-8'))
names=['BaoCao_Assignment2_PC-Astral-Market','UseCase_Logic_Ham_PC-Astral-Market']
records=[]
for i,name in enumerate(names):
    path=root/(name+'.docx');doc=Document(path)
    pdf=pdfium.PdfDocument(root/'tmp/documents-render'/(name+'.pdf'))
    pages=list((root/'tmp/documents-render'/name).glob('page-*.png'))
    assert len(pdf)==len(pages)
    assert len(doc.inline_shapes)==(state['figures'] if i==0 else 7)
    text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    for page in pdf:
        assert len(page.get_textpage().get_text_range().strip())>100, 'Trang trống hoặc gần trống'
    if i==0:
        assert state['headings']==[p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]
        assert len(doc.tables[0].rows)==4
        for row in doc.tables[0].rows[1:]:
            assert row.cells[1].text.strip()=='' and row.cells[2].text.strip()==''
        for phrase in ['Lê Anh Đào','Y1.7','Nhóm chọn tìm kiếm vì','LocalStorage','Product','Cart','GitHub','Không sử dụng thư viện bên ngoài']:
            assert phrase in text,phrase
        for old in ['Website chưa có tài khoản','sáu trang','Sáu trang','vào SessionStorage','fonts.css','components.css']:
            assert old not in text,old
        assert 'TOC ' in doc.element.xml
        toc=pdf[1].get_textpage().get_text_range()
        figures=pdf[2].get_textpage().get_text_range()
        assert len(re.findall(r'\.{3,}',toc))>=25
        assert len(re.findall(r'\.{3,}',figures))>=25
        assert 'Bảy trang' in figures
    else:
        filename=None;functions=0
        for block in doc.element.body:
            if block.tag.endswith('}p'):
                p=Paragraph(block,doc)
                if p.style.name=='Heading 2' and p.text.startswith('js/'):
                    filename=p.text
            elif block.tag.endswith('}tbl') and filename:
                table=Table(block,doc)
                if table.rows[0].cells[0].text!='Hàm hoặc phương thức':continue
                source=(root/'PC-Astral-Market'/filename).read_text(encoding='utf-8')
                for row in table.rows[1:]:
                    function=row.cells[0].text
                    match=re.search(r'(?m)^\s*(?:(?:async )?function )?'+re.escape(function)+r'\(',source)
                    assert match,(filename,function)
                    line=source[:match.end()].count('\n')+1
                    assert f'(dòng {line})' in row.cells[1].text,(filename,function,line)
                    functions+=1
        assert functions==45,functions
    records.append({'file':path.name,'pages':len(pdf),'figures':len(doc.inline_shapes),
        'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
        'all_pages_rendered_and_visually_reviewed':args.visual_reviewed})
allowed_documents = {'BaoCao_Assignment2_PC-Astral-Market.docx',
                     'UseCase_Logic_Ham_PC-Astral-Market.docx', 'Bộ-câu-hỏi.docx'}
assert {p.name for p in root.glob('*.docx')} == allowed_documents
summary={'date':date.today().isoformat(),'status':'PASS' if args.visual_reviewed else 'STRUCTURAL_PASS_PENDING_VISUAL',
    'renderer':'Microsoft Word COM PDF export and pypdfium2 PNG',
    'renderer_note':'render_docx.py không chạy được vì thiếu LibreOffice; dùng Word cài sẵn, giữ nguyên kiểu báo cáo.',
    'documents':records,'root_docx_count':3,'guide_functions_verified':45,
    'question_document_qa':'docs/js-revision-qa.json'}
(root/'docs/documents-qa.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(summary,ensure_ascii=True,indent=2))

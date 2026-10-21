"""Đóng gói website và báo cáo Y3 hiện hành; tài liệu tra cứu hàm để ngoài ZIP."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib, json, re, tempfile
from docx import Document
from sync_offline_products import sync_offline_products

# Mỗi lần đóng gói đều cập nhật bản mở trực tiếp từ JSON gốc, tránh dữ liệu lệch.
sync_offline_products()

root=Path(__file__).resolve().parents[1]
(root/'tmp').mkdir(exist_ok=True)
site=root/'PC-Astral-Market'
report=root/'BaoCao_Assignment2_PC-Astral-Market.docx'
archive=root/'submission/WEB1044_LapTrinhJSCoBan_PC-Astral-Market_Assignment.zip'
assert report.exists()
doc=Document(report)
assert len(doc.inline_shapes)==4
assert 'Lê Anh Đào' in '\n'.join(p.text for p in doc.paragraphs)
assert len(doc.tables[0].rows)==4
for row in doc.tables[0].rows[1:]: assert row.cells[1].text=='' and row.cells[2].text==''
assert all(p.is_relative_to(site/'js') for p in site.rglob('*.js'))
assert all(p.parent.name=='css' for p in site.rglob('*.css'))
assert len(list(site.glob('*.html')))==7
assert [p.name for p in (site/'css').glob('*.css')]==['style.css']
assert not list(site.rglob('*.svg'))
products=json.loads((site/'js/products.json').read_text(encoding='utf-8'))
assert len(products)==24 and len({p['category'] for p in products})==4
for product in products:
    assert all(key in product for key in ['id','name','price','imgUrl','description','category','quantity'])
    assert product['imgUrl'].startswith('img/') and (site/product['imgUrl']).is_file()
for p in site.rglob('*'):
    if p.is_file() and p.suffix in ['.html','.css','.js']:
        assert not re.search(r'[A-Z]:[\\/]', p.read_text(encoding='utf-8')), f'Absolute machine path in {p}'
with ZipFile(archive,'w',ZIP_DEFLATED) as z:
    for p in sorted(site.rglob('*')):
        if p.is_file(): z.write(p,p.relative_to(root).as_posix())
    z.write(report,report.name)
with ZipFile(archive) as z:
    assert z.testzip() is None
    destination=Path(tempfile.mkdtemp(prefix='submission-check-',dir=root/'tmp'))
    for member in z.infolist():
        assert (destination/member.filename).resolve().is_relative_to(destination.resolve())
    z.extractall(destination)
    assert report.read_bytes()==(destination/report.name).read_bytes()
    for p in site.rglob('*'):
        if p.is_file(): assert p.read_bytes()==(destination/p.relative_to(root)).read_bytes()
summary={'archive':archive.name,'entries':len(z.namelist()),'bytes':archive.stat().st_size,
         'sha256':hashlib.sha256(archive.read_bytes()).hexdigest(),'report':report.name,
         'report_figures':len(doc.inline_shapes),'members_left_blank':3,'website_pages':7,
         'extractedDirectory':str(destination)}
(root/'docs/package-check.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(summary,ensure_ascii=False,indent=2))

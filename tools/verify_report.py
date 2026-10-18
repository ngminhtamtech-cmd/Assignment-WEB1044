"""Check report fields, their actual PDF pages, and editable diagram structure."""
from pathlib import Path
from zipfile import ZipFile
from collections import Counter
from lxml import etree
from docx import Document
from pypdf import PdfReader
import hashlib
import json
import re
import argparse

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / 'submission/BaoCao_Assignment2_PC-Astral-Market.docx'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def normalize(text):
    return re.sub(r'\s+', ' ', text).strip()


def verify(baseline_path=None):
    manifest = json.loads((ROOT / 'docs/report-manifest.json').read_text(encoding='utf-8'))
    doc = Document(REPORT)
    pdf = PdfReader(ROOT / 'tmp/report/report.pdf')
    pages = [normalize(page.extract_text()) for page in pdf.pages]
    assert len(doc.inline_shapes) == manifest['figures'] == 19
    assert len(doc.tables) - 1 == manifest['tables'] == 6
    assert len(doc.tables[0].rows) == 4
    assert all(not cell.text for row in doc.tables[0].rows[1:] for cell in row.cells[1:])
    assert 'Lê Anh Đào' in '\n'.join(p.text for p in doc.paragraphs)
    assert abs(doc.sections[1].left_margin.cm - 3) < .01
    assert abs(doc.sections[1].right_margin.cm - 2) < .01
    assert not re.search(r'\bfunction\s*\(|\bconst\s+\w+\s*=|=>|\[mã code\]',
                         '\n'.join(p.text for p in doc.paragraphs))
    assert all(len(page) > 200 for page in pages), 'Unexpected blank or overflow page'

    with ZipFile(REPORT) as package:
        xml = etree.fromstring(package.read('word/document.xml'))
        styles = etree.fromstring(package.read('word/styles.xml'))
        relationships = etree.fromstring(package.read('word/_rels/document.xml.rels'))
    repository_url = 'https://github.com/ngminhtamtech-cmd/Assignment-WEB1044'
    assert any(r.get('Target') == repository_url and r.get('TargetMode') == 'External'
               for r in relationships), 'GitHub hyperlink is missing'
    assert any(repository_url in page for page in pages), 'GitHub URL is not visible'
    style_map = {s.get(W + 'styleId'): s for s in styles}
    def font_size(style):
        value = style.find('./' + W + 'rPr/' + W + 'sz')
        if value is not None:
            return value.get(W + 'val')
        base = style.find(W + 'basedOn')
        if base is not None:
            return font_size(style_map[base.get(W + 'val')])
        return styles.find('./' + W + 'docDefaults/' + W + 'rPrDefault/' + W + 'rPr/' + W + 'sz').get(W + 'val')
    toc_styles = {'TOC1', 'TOC2', 'TableofFigures'}
    for name in toc_styles:
        style = style_map[name]
        tabs = style.findall('./' + W + 'pPr/' + W + 'tabs/' + W + 'tab')
        assert any(t.get(W + 'val') == 'right' and t.get(W + 'leader') == 'dot'
                   and abs(int(t.get(W + 'pos')) - 9071) <= 2 for t in tabs), name
        assert font_size(style) == '22'
    level_two = style_map['TOC2'].find('./' + W + 'pPr/' + W + 'ind')
    assert abs(int(level_two.get(W + 'left')) - 340) <= 2
    instructions = [item.text or '' for item in xml.iter(W + 'instrText')]
    assert any('TOC \\o "1-2"' in instruction for instruction in instructions)

    entries = []
    heading_titles = []
    caption_numbers = {'Hình': [], 'Bảng': []}
    for p in xml.iter(W + 'p'):
        st = p.find('./' + W + 'pPr/' + W + 'pStyle')
        style = st.get(W + 'val') if st is not None else ''
        if style in {'Heading1', 'Heading2'}:
            heading_titles.append(normalize(''.join(p.itertext(W + 't'))))
        if style == 'Caption':
            text = ''.join(p.itertext(W + 't'))
            match = re.match(r'(Hình|Bảng)\s+(\d+)\s', text)
            assert match, text
            caption_numbers[match[1]].append(int(match[2]))
        if style not in toc_styles:
            continue
        link = p.find(W + 'hyperlink')
        assert link is not None
        pieces = [t.text or '' for t in link.iter(W + 't')]
        title, page = normalize(''.join(pieces[:-1])), int(pieces[-1])
        assert len(list(link.iter(W + 'tab'))) == 1
        assert '...' not in title, 'Leaders must be a tab setting, not text'
        assert title in pages[page], (title, page)
        anchor = link.get(W + 'anchor')
        assert xml.find('.//' + W + 'bookmarkStart[@' + W + 'name="' + anchor + '"]') is not None
        entries.append({'title': title, 'page': page, 'style': style})
    toc_titles = [entry['title'] for entry in entries if entry['style'] != 'TableofFigures']
    assert toc_titles == heading_titles
    assert not any(title in toc_titles for title in ['Mục lục', 'Danh mục hình', 'Danh mục bảng'])
    assert caption_numbers['Hình'] == list(range(1, 20))
    assert caption_numbers['Bảng'] == list(range(1, 7))

    diagram = etree.parse(str(ROOT / 'docs/diagrams/usecases/PC-Astral-Market-UseCases.drawio'))
    drawing_pages = diagram.findall('diagram')
    assert len(drawing_pages) == 3
    expected_ids = [None, {f'UC{i:02}' for i in range(1, 7)}, {f'UC{i:02}' for i in range(6, 12)}]
    expected_edges = [5, 6, 9]
    for index, page in enumerate(drawing_pages):
        cells = page.findall('.//mxCell')
        ids = [cell.get('id') for cell in cells]
        assert len(ids) == len(set(ids))
        assert len([cell for cell in cells if 'umlActor' in cell.get('style', '')]) == 1
        usecases = [cell for cell in cells if cell.get('style', '').find('ellipse;') >= 0]
        assert len(usecases) == [5, 6, 8][index]
        if expected_ids[index]:
            assert {cell.get('id') for cell in usecases if cell.get('id').startswith('UC')} == expected_ids[index]
        edges = [cell for cell in cells if cell.get('edge') == '1']
        assert len(edges) == expected_edges[index]
        for edge in edges:
            assert edge.get('source') in ids and edge.get('target') in ids
        if index == 2:
            includes = {(e.get('source'), e.get('target')) for e in edges if 'dashed=1' in e.get('style', '')}
            assert includes == {('UC06', 'check-quantity'), ('UC08', 'check-quantity'), ('UC10', 'check-customer')}
    actual = {p.relative_to(ROOT / 'PC-Astral-Market').as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
              for p in (ROOT / 'PC-Astral-Market').rglob('*') if p.is_file()}
    if baseline_path is not None:
        baseline = json.loads(Path(baseline_path).read_text(encoding='utf-8'))
        assert baseline == actual, 'Website files must remain unchanged'
    result = {
        'status': 'PASS', 'pages': len(pages), 'figures': 19, 'tables': 6,
        'toc_and_list_entries_checked': len(entries), 'index_counts': dict(Counter(e['style'] for e in entries)),
        'right_tab_cm': 16, 'leader': 'dot', 'font': 'Times New Roman 11',
        'all_index_pages_match_pdf': True, 'drawio_pages': 3, 'include_relations': 3,
        'website_files': len(actual), 'website_baseline_checked': baseline_path is not None,
        'github_hyperlink': repository_url, 'member_rows_blank': 3,
        'report_sha256': hashlib.sha256(REPORT.read_bytes()).hexdigest(), 'entries': entries,
    }
    (ROOT / 'docs/report-validation.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: v for k, v in result.items() if k != 'entries'}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--site-baseline', type=Path, help='Optional JSON of website file hashes before the report edit')
    verify(parser.parse_args().site_baseline)

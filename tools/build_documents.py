"""Chỉ dựng tài liệu Use Case, Logic và Hàm; tuyệt đối không ghi đè báo cáo Y3."""
from pathlib import Path
import re, json
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from report_usecases import make_usecases

ROOT = Path(__file__).resolve().parents[1]
SITE = ROOT/'PC-Astral-Market'
FIG = ROOT/'docs/diagrams/concise'
FIG.mkdir(parents=True, exist_ok=True)
USE = make_usecases()
GUIDE = ROOT/'UseCase_Logic_Ham_PC-Astral-Market.docx'
REPO = 'https://github.com/ngminhtamtech-cmd/Assignment-WEB1044'

def flow(name, labels, note):
    image = Image.new('RGB', (1600, 360), 'white')
    d = ImageDraw.Draw(image)
    font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 29)
    step = 1600/len(labels)
    for i, label in enumerate(labels):
        x = int(i*step)+12
        d.rounded_rectangle((x,30,x+step-45,195),radius=10,outline='#333333',width=3)
        lines=label.split('\n')
        for j,line in enumerate(lines):
            assert d.textlength(line,font=font)<step-65, line
            d.text((x+(step-45)/2,70+j*38),line,font=font,fill='#222222',anchor='mm')
        if i<len(labels)-1:
            y=112; end=x+step-4
            d.line((x+step-45,y,end,y),fill='#333333',width=3)
            d.polygon(((end,y),(end-13,y-8),(end-13,y+8)),fill='#333333')
    d.multiline_text((20,230),note,font=font,fill='#222222',spacing=12)
    image.save(FIG/(name+'.png'))
    return FIG/(name+'.png')

CART = flow('gio-hang', ['Thêm / sửa / xóa\nKiểm tra số lượng','Cập nhật Cart\nGộp dòng trùng','Lưu LocalStorage\nTính lại tổng','Cập nhật giao diện\nDòng giỏ và badge'],
            'F5: đọc dữ liệu đã lưu → kiểm tra dòng giỏ → tạo lại Cart → hiển thị.\nNhập sai: báo lỗi và giữ giá trị hợp lệ trước đó. Không sửa riêng con số trên màn hình.')
ORDER = flow('dat-hang', ['Kiểm tra form\nGiỏ và giá','Tạo bản chụp đơn\nLưu đơn thành công','Xóa nháp\nDọn giỏ khớp đơn','Hiển thị xác nhận\nGiữ nguyên mã đơn'],
             'Lưu đơn lỗi: giữ form và giỏ. Giá hoặc giỏ đổi: yêu cầu xác nhận lại.\nDọn giỏ lỗi: đơn vẫn đã lưu; nút thử lại chỉ dọn giỏ, không tạo thêm đơn.')
SLIDER = flow('slider', ['Nhận nút / timer\nTrước, sau, chấm','Đổi vị trí banner\nVòng đầu ↔ cuối','Đọc banner tại vị trí\nẢnh, tiêu đề, link','Cập nhật giao diện\nĐánh dấu chấm chọn'],
              'Tự chuyển sau 6 giây; tạm dừng khi rê chuột, có focus, tab ẩn hoặc người dùng bấm dừng.')
CLOCK = flow('countdown', ['Đọc hạn cố định\nTrong du-lieu.js','Hạn − giờ hiện tại\nKhông nhỏ hơn 0','Đổi ra ngày giờ\nPhút và giây','Hết hạn: dừng đếm\nTrở về giá thường'],
             'Cập nhật mỗi giây và khi trở lại tab. F5 không tạo lại thời hạn.\nProduct.layGia là nguồn đơn giá chung của danh sách, chi tiết, giỏ và đơn hàng.')

def setup():
    doc=Document(); s=doc.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.left_margin=Cm(3);s.right_margin=Cm(2);s.top_margin=s.bottom_margin=Cm(2)
    s.header_distance=s.footer_distance=Cm(.8)
    for name,size in [('Normal',12),('Title',21),('Subtitle',12),('Heading 1',15),('Heading 2',13),('Caption',10.5),('TOC 1',11),('TOC 2',11)]:
        if name not in doc.styles: doc.styles.add_style(name,1)
        st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        for border in st.element.xpath('.//w:pBdr'):border.getparent().remove(border)
        for fonts in st.element.xpath('.//w:rFonts'):
            for attribute in list(fonts.attrib):
                if 'theme' in attribute.lower():del fonts.attrib[attribute]
            for attribute in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+attribute),'Times New Roman')
        st.paragraph_format.line_spacing=1.15;st.paragraph_format.space_after=Pt(7)
        if name.startswith('Heading'):st.font.bold=True;st.paragraph_format.keep_with_next=True
        if name.startswith('TOC'):
            st.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
            st.paragraph_format.space_after=Pt(5)
    foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    return doc

def para(doc,text,style=None): return doc.add_paragraph(text,style)
def heading(doc,text): return doc.add_heading(text,1)
def page(doc,title): doc.add_page_break();heading(doc,title)
def table(doc,headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Cm(w)
    for i,text in enumerate(headers): t.rows[0].cells[i].text=text
    for row in rows:
        cells=t.add_row().cells
        for i,text in enumerate(row):cells[i].text=str(text)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr()
        keep=OxmlElement('w:cantSplit');pr.append(keep)
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for i,c in enumerate(row.cells):
            c.width=Cm(widths[i]);tc=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tc.append(borders)
            margins=OxmlElement('w:tcMar')
            for side,value in [('top','70'),('bottom','70'),('left','90'),('right','90')]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');margins.append(e)
            tc.append(margins)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF2');tc.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05
                for r in p.runs:r.font.size=Pt(11);r.bold=ri==0
    return t

def figure(doc,path,caption,width=16):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Cm(width))
    p=para(doc,caption,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER

def link(doc,label,url):
    p=doc.add_paragraph();h=OxmlElement('w:hyperlink')
    h.set(qn('r:id'),doc.part.relate_to(url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def toc(doc):
    p=doc.add_paragraph('MỤC LỤC');p.runs[0].bold=True
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z \\u')
    p=doc.add_paragraph();p._p.append(f)

def funcs(doc,filename,entries):
    source=(SITE/filename).read_text(encoding='utf-8')
    doc.add_heading(filename,2)
    rows=[]
    for name,description in entries:
        m=re.search(r'(?m)^\s*(?:(?:async )?function )?'+re.escape(name)+r'\(',source)
        assert m, (filename,name)
        lineno=source[:m.end()].count('\n')+1
        rows.append((name,description+' (dòng '+str(lineno)+').'))
    table(doc,['Hàm hoặc phương thức','Vai trò và vị trí'],rows,[5.4,10.6])

doc=setup()
para(doc,'PC ASTRAL MARKET','Title')
para(doc,'Use Case logic hoạt động và các hàm','Subtitle')
heading(doc,'1 Use Case tổng quan')
para(doc,'Tài liệu này dùng để lần theo chức năng vào mã nguồn. Đọc sơ đồ để biết khách thực hiện được gì, đọc luồng để biết dữ liệu thay đổi thế nào, rồi dùng bảng hàm để tìm đúng chỗ sửa. Tên hàm và số dòng được lấy theo mã hiện tại; số dòng có thể thay đổi khi sửa file.')
figure(doc,USE/'01-overview.png','Hình 1. Phạm vi sử dụng của khách hàng',11.4)
para(doc,'Khách không cần đăng nhập để đặt hàng. Tài khoản demo hỗ trợ đăng ký, đăng nhập và đăng xuất. “Kết thúc buổi demo” là thao tác xóa dữ liệu riêng của website; không phải chức năng quản trị. Website không có tài khoản quản trị hay thanh toán trực tuyến.')

page(doc,'2 Use Case tìm và chọn sản phẩm')
figure(doc,USE/'02-find-select.png','Hình 2. Các chức năng tra cứu và chọn hàng',12.0)
para(doc,'Mục tiêu: tìm sản phẩm phù hợp. Điều kiện trước: dữ liệu sản phẩm đã nạp. Khách có thể xem danh mục, tìm kiếm, lọc và sắp xếp riêng hoặc kết hợp; sau đó xem chi tiết hoặc thêm giỏ ngay từ danh sách.')
para(doc,'Khi không có kết quả, trang hiện thông báo và cho xóa lọc. ID không tồn tại hiện “Không tìm thấy sản phẩm”. Sản phẩm hết hàng không cho thêm. Kết quả: danh sách đúng điều kiện và giỏ nhận đúng mã hàng, số lượng. UC02–UC04 là ba yêu cầu tự đề xuất của Y1.7.')

page(doc,'3 Use Case giỏ hàng và đặt hàng')
figure(doc,USE/'03-cart-checkout.png','Hình 3. Giỏ hàng và các bước kiểm tra bắt buộc',12.0)
para(doc,'Mục tiêu: hoàn tất một đơn mô phỏng. Khách xem giỏ, sửa hoặc xóa hàng; số lượng chỉ nhận số nguyên từ 1 đến tồn kho. Thêm trùng tăng số lượng của dòng cũ. Giỏ rỗng không thể đặt đơn.')
para(doc,'Đặt hàng luôn kiểm tra thông tin người nhận. Mũi tên «include» đi từ chức năng chính tới bước kiểm tra được sử dụng. Kết quả thành công là một bản chụp đơn đã lưu, không phải giao dịch thật. “Xem xác nhận” là đơn gần nhất được lưu trên trình duyệt, không phải danh sách lịch sử đơn hàng.')

page(doc,'4 Logic giỏ hàng đặt hàng và F5')
figure(doc,CART,'Hình 4. Dữ liệu giỏ là nguồn tính tổng')
para(doc,'Cart giữ các dòng gồm productId và quantity. Khi tính tiền, Cart tra Product theo ID, lấy giá đang áp dụng, nhân số lượng rồi cộng dồn. Định dạng tiền chỉ thực hiện khi hiển thị, vì chuỗi “1.000.000 đ” không dùng để cộng tiền.')
para(doc,'Nút trong danh sách và giỏ được xử lý bằng event delegation ở phần tử cha. Sự kiện xác định nút gần nhất và mã sản phẩm, gọi Cart trước rồi cập nhật dòng hiển thị. Nhập số lượng sai giữ số cũ, không dựng lại toàn bộ form làm mất focus.')
figure(doc,ORDER,'Hình 5. Chỉ dọn giỏ sau khi đơn lưu được')
para(doc,'Khi form đang mở, ứng dụng so cả mã hàng, số lượng và đơn giá với lần hiển thị trước. Dù tổng tiền không đổi, đổi sản phẩm vẫn phải xác nhận lại. Form sai hoặc lưu đơn thất bại đều giữ dữ liệu để sửa hoặc thử lại.')
para(doc,'Đơn đã lưu nhưng dọn giỏ lỗi: cờ choDonGio cho biết bước chưa xong. F5 đọc lại đơn; nút thử lại chỉ dọn khi giỏ còn khớp đơn. Nếu khách đã mua giỏ mới, không tự xóa giỏ đó. Khóa trạng thái gửi ngăn nhấp nhanh hoặc Enter tạo nhiều đơn.')
para(doc,'Bản nháp lưu đủ năm trường mỗi khi nhập và được khôi phục khi F5. Dữ liệu nháp hỏng được bỏ qua; không lưu mật khẩu đang gõ. Form đang nhập không bị dữ liệu nháp của tab khác tự ghi đè.')

page(doc,'5 Logic slider countdown và tài khoản')
figure(doc,SLIDER,'Hình 6. Chuyển banner bằng JavaScript')
figure(doc,CLOCK,'Hình 7. Countdown dùng thời hạn cố định')
para(doc,'Đăng ký: kiểm tra tên và mật khẩu → tìm tên trùng → tạo salt → băm mật khẩu → kiểm tra dấu reset → lưu username, salt và hash. Đăng nhập: tìm tài khoản → băm mật khẩu nhập với salt đã lưu → so hash → kiểm tra reset và tài khoản còn tồn tại → lưu trạng thái đăng nhập.')
para(doc,'Phần băm phải chờ xử lý bất đồng bộ. Nếu người dùng kết thúc demo trong lúc chờ, mã kiểm tra dấu reset sau await và dừng, tránh khôi phục dữ liệu vừa bị xóa. Không bỏ bước này để làm mã ngắn hơn.')
table(doc,['Khóa LocalStorage','Dữ liệu và lúc xóa'],[
('pc-astral.cart.v1','Giỏ; dọn sau khi lưu đơn thành công hoặc kết thúc demo.'),
('pc-astral.checkout-draft.v1','Nháp năm trường; xóa sau lưu đơn hoặc kết thúc demo.'),
('pc-astral.users.v1','Username, salt, hash; giữ sau đăng xuất, chỉ website chủ động xóa khi kết thúc demo.'),
('pc-astral.auth.v1','Trạng thái đăng nhập; xóa khi đăng xuất hoặc kết thúc demo.'),
('pc-astral.last-order.v1','Đơn gần nhất; đơn mới thay thế đơn cũ, kết thúc demo thì xóa.'),
('pc-astral.demo-reset.v1','Mốc reset để chặn tác vụ cũ; được cập nhật và giữ lại.')],[6.5,9.5])

page(doc,'6 Các hàm dữ liệu và mô hình')
funcs(doc,'js/du-lieu.js',[
('kiemTraSanPham','Kiểm tra kiểu và các trường bắt buộc trước khi tạo Product'),
('napDuLieuSanPham','Đọc JSON qua HTTP; mở file trực tiếp thì đọc bản nhúng')])
funcs(doc,'js/mo-hinh.js',[
('dangGiamGia','So chiến dịch và thời hạn; chỉ giảm khi giá giảm hợp lệ'),
('layGia','Trả về số tiền hiện hành dùng chung cho mọi trang'),
('layGiaHienThi','Đổi số thành chuỗi VNĐ; không thay đổi giá gốc'),
('conHang','Cho biết tồn kho còn lớn hơn 0'),
('coTheMua','Kiểm tra số lượng nguyên an toàn và không vượt tồn kho'),
('layTenDanhMuc','Tra tên danh mục theo mã'),
('kiemTraGio','Lọc dữ liệu lưu bị hỏng, gộp trùng và giới hạn tồn kho'),
('themSanPham','Tìm dòng cũ để cộng dồn hoặc tạo dòng mới'),
('doiSoLuong','Kiểm tra số mới, cập nhật dòng và gọi lưu giỏ'),
('xoaSanPham','Xóa dòng theo ID rồi báo dữ liệu thay đổi'),
('tinhTongGio','Tính đơn giá, thành tiền, tổng số lượng và tổng tiền'),
('layTongTien','Lấy tổng tiền từ kết quả tính giỏ'),
('xoaGio','Thử lưu giỏ rỗng; phục hồi giỏ trong bộ nhớ nếu lưu thất bại')])
para(doc,'Product có thuộc tính sản phẩm; Cart giữ danh sách sản phẩm, các dòng giỏ và hàm gọi khi thay đổi. constructor là tên chuẩn của JavaScript. Dùng Ctrl+F tên hàm trong đúng file; không đổi tên để tránh đứt các nơi đang gọi.')

page(doc,'7 Các hàm lưu trữ và đặt hàng')
funcs(doc,'js/luu-tru.js',[
('docJSON','Đọc và JSON.parse; dữ liệu lỗi trả về giá trị dự phòng'),
('luuJSON','JSON.stringify rồi ghi; trả true/false để bên gọi xử lý'),
('ketThucDemo','Ghi mốc reset rồi xóa đúng các khóa dữ liệu website'),
('bamMatKhau','Dùng Web Crypto PBKDF2 với salt, không lưu mật khẩu thô'),
('dangKy','Kiểm tra, băm và lưu tài khoản; chặn tên trùng'),
('dangNhap','So hash, kiểm tra reset sau await rồi lưu phiên'),
('dangXuat','Xóa trạng thái đăng nhập, giữ tài khoản và giỏ')])
funcs(doc,'js/trang/dat-hang.js',[
('luuBanNhap','Lưu năm trường khi form phát sinh input'),
('khoiPhucBanNhap','Đọc nháp và điền lại các trường có kiểu chuỗi'),
('cacDongGiongNhau','So ID, số lượng và khi cần so thêm giá'),
('hienThiTongTien','Cập nhật tóm tắt; đánh dấu khi cần xác nhận lại'),
('taoDonHang','Tạo mã đơn, thời điểm, thông tin và bản chụp dòng hàng'),
('xuLyDatHang','Kiểm tra form/giá, lưu đơn trước rồi xử lý dọn giỏ'),
('thuDonGio','Thử lại bước dọn giỏ của cùng đơn; không tạo đơn khác'),
('hienThiXacNhan','Hiển thị bản chụp đã lưu và cảnh báo còn tồn tại')])

page(doc,'8 Các hàm hiển thị và sự kiện')
funcs(doc,'js/hien-thi.js',[
('boDau','Chuẩn hóa tìm kiếm không dấu và không phân biệt hoa thường'),
('locSanPham','Duyệt sản phẩm, kết hợp tìm kiếm và các điều kiện lọc'),
('kiemTraKhachHang','Trim rồi kiểm tra bắt buộc, email và điện thoại'),
('taoTheSanPham','Dựng thẻ dùng chung từ Product'),
('theoDoiKhuyenMai','Phát sự kiện giá cần cập nhật khi hết hạn')])
funcs(doc,'js/su-kien.js',[
('xuLyThemGio','Bắt nút thêm giỏ qua phần tử cha và gọi Cart'),
('xuLyLuuTruThayDoi','Đồng bộ khi tab khác thay đổi dữ liệu'),
('khoiTaoWebsite','Chờ dữ liệu rồi tạo mô hình, giao diện và trang')])
funcs(doc,'js/trang/trang-chu.js',[
('chuyenBanner','Cộng vị trí, vòng đầu/cuối, hiển thị banner mới'),
('tuChuyenBanner','Kiểm tra điều kiện tạm dừng trước khi tự chuyển'),
('capNhatDemNguoc','Tính phần thời gian còn lại, dừng khi bằng 0')])
funcs(doc,'js/trang/san-pham.js',[
('hienThiDanhSachSanPham','Lấy kết quả lọc, sắp xếp rồi render'),
('capNhatURL','Ghi lựa chọn vào URL để reload vẫn giữ')])
funcs(doc,'js/trang/gio-hang.js',[
('xuLyBamNutGio','Nhận tăng/giảm/xóa qua event delegation'),
('xuLyDoiSoLuong','Đọc ô nhập và giữ giá trị trước nếu nhập sai')])
doc.save(GUIDE)
print(GUIDE.name)

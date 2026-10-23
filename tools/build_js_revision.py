"""Bản ôn JS ngắn thay thế Bộ-câu-hỏi; không thay website hay Word khác."""
from pathlib import Path
import json, hashlib, re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from defense_content import QUESTIONS

ROOT=Path(__file__).resolve().parents[1]
SITE=ROOT/'PC-Astral-Market'
data={q['id']:q for q in QUESTIONS}
groups=[
('Biến hàm và class',['B01','B02','A08'],'let có thể gán lại; const không gán lại biến. Hàm nhận tham số, xử lý và return kết quả. Product giữ sản phẩm; Cart giữ giỏ.', 'Trong constructor, this là đối tượng đang tạo bằng new. JSON.parse chỉ tạo dữ liệu thường; muốn có phương thức thì tạo lại Product.'),
('Đọc dữ liệu và khởi tạo',['A05','A06','B06'],'JSON → tạo Product → tạo Cart từ dữ liệu lưu → dựng giao diện → gọi logic trang.', 'defer giữ thứ tự script; await chờ đọc JSON. window.PCAstral là object chung để các file trao đổi dữ liệu và hàm.'),
('DOM và sự kiện',['A09','A10','A11'],'Dữ liệu → tạo thẻ → đưa vào DOM. Người dùng bấm → tìm nút → lấy ID → gọi xử lý.', 'addEventListener gắn sự kiện. event.target là nơi phát sinh sự kiện; closest tìm phần tử khớp gần nhất, gồm cả chính phần tử đó. Không gắn listener lại trong mỗi lần render.'),
('Thêm sửa và tính tổng giỏ',['A12','A13','A14'],'Tìm ID → kiểm tra số lượng → sửa Cart → lưu → cập nhật giao diện.', 'Tổng = cộng từng đơn giá × số lượng bằng for. Sửa dữ liệu trước, không chỉ sửa chữ trên DOM. Nhập số sai giữ số lượng cũ; muốn bỏ dòng dùng Xóa.'),
('LocalStorage và F5',['A15','A16','B16'],'Lưu: dữ liệu → stringify → setItem. Đọc: getItem → parse → kiểm tra → tạo Cart.', 'F5 không tự xóa LocalStorage. Dữ liệu phụ thuộc origin và hồ sơ trình duyệt. Khi storage bị chặn, trang có thể thao tác trong bộ nhớ nhưng không bảo đảm giữ sau F5.'),
('Đăng ký đăng nhập và reset',['A17','A18','A19'],'Đăng ký → kiểm tra → tạo salt → băm → lưu. Đăng nhập → đọc tài khoản → băm mật khẩu nhập → so sánh.', 'Không lưu mật khẩu nguyên văn. async/await chờ băm; sau await kiểm tra dấu reset để không ghi lại dữ liệu vừa xóa. Đăng nhập phía trình duyệt chỉ là mô phỏng, không thay backend.'),
('Tìm kiếm lọc và sắp xếp',['B07','B08','B09'],'Chuẩn hóa từ khóa → for lọc theo AND → sort kết quả → render.', 'Tìm kiếm bỏ dấu và đổi về chữ thường. sort thay đổi mảng được gọi, nên chỉ sort mảng kết quả riêng. Bộ lọc được ghi vào URL và đọc lại khi F5.'),
('Slider countdown và giá',['B11','B13','B14'],'Slider đổi vị trí → kiểm tra đầu/cuối → hiển thị. Countdown lấy hạn cố định trừ giờ hiện tại.', 'setInterval gọi định kỳ; slider tạm dừng khi hover, focus, tab ẩn hoặc bấm dừng. Countdown chặn thời gian âm; giá dùng chung Product.layGia để các trang nhất quán.'),
('Form và đặt hàng',['B26','A20','B28'],'Kiểm tra form → kiểm tra giỏ/giá → lưu đơn thành công → xóa nháp → dọn giỏ → xác nhận.', 'Giá hoặc giỏ thay đổi phải yêu cầu xác nhận lại, giữ form. Khóa trạng thái gửi để tránh lặp; F5 trang xác nhận đọc đơn đã lưu, không tạo đơn mới.')]

d=Document();s=d.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7)
s.left_margin=Cm(3);s.right_margin=s.top_margin=s.bottom_margin=Cm(2)
for name,size in [('Normal',12),('Title',20),('Heading 1',16),('Heading 2',12),('TOC 1',12)]:
 if name not in d.styles:d.styles.add_style(name,1)
 st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
 fonts=st.element.get_or_add_rPr().find(qn('w:rFonts'))
 for key in list(fonts.attrib):
  if 'theme' in key.lower():del fonts.attrib[key]
 for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),'Times New Roman')
 st.paragraph_format.line_spacing=1.15;st.paragraph_format.space_after=Pt(6)
 if name.startswith('Heading'):st.font.bold=True;st.paragraph_format.keep_with_next=True
 if name=='TOC 1':
  st.element.attrib.pop(qn('w:customStyle'),None);st.element.find(qn('w:name')).set(qn('w:val'),'toc 1')
  st.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
for st in d.styles:
 for b in st.element.xpath('.//w:pBdr'):b.getparent().remove(b)
s.header.paragraphs[0].text='PC ASTRAL MARKET  |  ÔN JAVASCRIPT'
s.header.paragraphs[0].runs[0].font.size=Pt(9)
f=s.footer.paragraphs[0];f.alignment=2
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
important=[
 'không cho gán lại chính biến', 'không tự động đóng băng',
 'chỉ dùng được trong phạm vi tương ứng', 'không tự có phương thức',
 'chỉ là dữ liệu thường', 'ID riêng', 'đọc và kiểm tra JSON',
 'theo thứ tự trong tài liệu', 'không có nghĩa dữ liệu đã về',
 'không cần gắn lại từng nút', 'không gắn listener lại',
 'cộng số lượng mới vào dòng cũ', 'không vượt tồn kho',
 'đơn giá × số lượng', 'giữ số lượng cũ', 'trở lại số lượng hợp lệ trước đó',
 'số nguyên từ 1 đến tồn kho', 'Sửa dữ liệu trước',
 'không tự xóa LocalStorage', 'không tự xóa tài khoản',
 'không bảo đảm giữ sau F5', 'không lưu mật khẩu nguyên văn',
 'Không lưu mật khẩu nguyên văn', 'băm mật khẩu vừa nhập',
 'so sánh với passwordHash đã lưu', 'chỉ xóa khóa trạng thái đăng nhập',
 'xóa riêng các khóa ứng dụng', 'kiểm tra dấu reset',
 'chỉ là mô phỏng', 'tắt server không tự kết thúc demo',
 'thỏa tất cả điều kiện', 'không đổi', 'mảng kết quả riêng',
 'chưa có xác nhận thành công', 'lưu thành công rồi mới xóa',
 'giữ form và giỏ', 'chỉ cho thử dọn lại', 'không tạo đơn mới',
 'yêu cầu xác nhận lại', 'giữ form', 'Hạn kết thúc được lưu cố định',
 'chặn thời gian âm', 'giá dùng chung', 'lưu nguyên năm trường'
]
tokens=['let','const','return','for','if','continue','new','this','constructor','class',
 'Product','Cart','JSON','LocalStorage','DOM','ID','F5','AND','URL','defer','await',
 'async/await','try/catch','sort','includes','stringify','parse','setItem','getItem',
 'setInterval','clearInterval','addEventListener','closest','target','currentTarget',
 'username','salt','passwordHash','quantity','price','salePrice','textContent','innerHTML']
emphasis=re.compile('|'.join(re.escape(x) for x in sorted(important,key=len,reverse=True))+
 r'|\b(?:window\.PCAstral|JSON\.(?:parse|stringify)|event\.target|Date\.now|(?:Product|Cart)\.[A-Za-z]+|[a-z]+(?:[A-Z][a-zA-Z0-9]*)+)\b|\b(?:'+ '|'.join(re.escape(x) for x in tokens)+r')\b')
def p(text,label=None):
 x=d.add_paragraph()
 if label:x.add_run(label+': ').bold=True
 cursor=0
 for match in emphasis.finditer(text):
  x.add_run(text[cursor:match.start()])
  x.add_run(match.group()).bold=True
  cursor=match.end()
 x.add_run(text[cursor:])
 return x

d.add_paragraph('ÔN JAVASCRIPT\nPC ASTRAL MARKET',style='Title')
p('Bản tóm tắt chỉ tập trung vào code JS của dự án: 27 câu hỏi ngắn trong 9 phần. Đọc để hiểu luồng, mở đúng hàm và tự giải thích; không cần học thuộc nguyên văn.')
p('Nhận gì → kiểm tra gì → xử lý gì → trả về hoặc cập nhật gì.', 'Cách giải thích một hàm')
p('Mở file bằng Ctrl+P; đến dòng bằng Ctrl+G. Đường dẫn tính từ PC-Astral-Market. Số dòng đối chiếu ngày 23/10/2026; nếu mã đã đổi, tìm theo tên hàm.')
p('Ưu tiên phần 4, 5, 6 và 9: giỏ hàng, lưu dữ liệu, tài khoản và đặt hàng.', 'Ôn nhanh')
p('NỘI DUNG').runs[0].bold=True
toc=OxmlElement('w:fldSimple');toc.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z');d.add_paragraph()._p.append(toc)
p('Ví dụ dễ nhầm: F5 không tự xóa tài khoản; tổng tiền đang dùng for; mật khẩu lưu dạng mã băm; tắt server không tự kết thúc demo.', 'Nhớ đúng')

symbols=json.loads((ROOT/'docs/defense-source-symbols.json').read_text(encoding='utf-8'))
refs=[]
for i,(title,ids,flow,note) in enumerate(groups,1):
 d.add_page_break();d.add_heading(f'{i} {title}',1)
 p(flow,'Luồng cần nhớ')
 for oldid in ids:
  q=data[oldid];d.add_heading(q['ask'],2)
  answer=q['answer']
  p(answer)
  links=[]
  for file,anchor,label in q['refs']:
   if file.startswith('tools/'):continue
   lines=(SITE/file).read_text(encoding='utf-8').splitlines()
   line=next(n+1 for n,t in enumerate(lines) if anchor in t)
   candidates=[f for f in symbols.get(file,[]) if f['start']<=line<=f['end']]
   func=min(candidates,key=lambda f:f['end']-f['start']) if candidates and not anchor.strip().startswith('class ') else None
   name=func['name'] if func else label
   links.append(f'{file} — {name}, dòng {line}')
   refs.append(dict(file=file,anchor=anchor,line=line))
  p('; '.join(links)+'.','Mở code')
 p(note,'Chốt lại')
 p('Mở các hàm trên, đọc đầu vào và return; thử một trường hợp đúng và một trường hợp sai. Không sửa bản nộp để luyện lỗi.','Tự kiểm tra')

d.save(ROOT/'Bộ-câu-hỏi.docx')
(ROOT/'docs/js-revision-index.json').write_text(json.dumps({'topics':9,'questions':27,'references':refs},ensure_ascii=False,indent=2),encoding='utf-8')
print('Created JS-only revision with 27 questions and',len(refs),'source references')

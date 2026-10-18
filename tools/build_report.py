from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from report_visuals import make_diagrams, make_cover_frame, OUT as FIGURES
from report_usecases import make_usecases
import json

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'submission'
QA=ROOT/'tmp/report'
OUT.mkdir(exist_ok=True)
QA.mkdir(parents=True,exist_ok=True)
make_diagrams()
USECASES=make_usecases()
frame_path=make_cover_frame()
doc=Document()
sec=doc.sections[0]
sec.page_width,sec.page_height=Cm(21),Cm(29.7)
sec.left_margin=sec.right_margin=Cm(2.5)
sec.top_margin=sec.bottom_margin=Cm(2)
sec.header_distance=sec.footer_distance=Cm(.8)
sec.different_first_page_header_footer=True
for name,size in [('Normal',13),('Title',20),('Subtitle',14),('Heading 1',17),('Heading 2',14),('Heading 3',13),('Caption',11)]:
    st=doc.styles[name]
    st.font.name='Times New Roman'
    st.font.size=Pt(size)
    st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.line_spacing=1.5
    if name.startswith('Heading'):
        st.font.bold=True
        st.paragraph_format.space_before=Pt(8)
        st.paragraph_format.space_after=Pt(8)
        st.paragraph_format.keep_with_next=True
    if name=='Caption':
        st.font.italic=True
        st.paragraph_format.line_spacing=1.1
        st.paragraph_format.space_after=Pt(8)
    if name=='Title':
        st.font.bold=True
        st.paragraph_format.line_spacing=1.3
for name in ['TOC 1','TOC 2','Table of Figures']:
    if name not in doc.styles:
        doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH,builtin=True)
    st=doc.styles[name]
    st.font.name='Times New Roman';st.font.size=Pt(11)
    st.font.color.rgb=RGBColor(0,0,0)
    st.font.bold=name=='TOC 1'
    st.paragraph_format.line_spacing=1;st.paragraph_format.space_after=Pt(4)
    st.paragraph_format.space_before=Pt(0)
    st.paragraph_format.left_indent=Cm(.6 if name=='TOC 2' else 0)
    st.paragraph_format.first_line_indent=Cm(0)
    st.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
st=doc.styles.add_style('Front Matter Title',WD_STYLE_TYPE.PARAGRAPH)
st.base_style=doc.styles['Normal'];st.font.bold=True;st.font.size=Pt(14)
st.paragraph_format.keep_with_next=True;st.paragraph_format.space_after=Pt(8)
outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),'9');st.element.get_or_add_pPr().append(outline)
doc.core_properties.title='Báo cáo Assignment 2 PC Astral Market'
doc.core_properties.subject='WEB1044 Lập trình cơ sở với JavaScript'
doc.core_properties.author='Nhóm sinh viên'
doc.core_properties.keywords='WEB1044, Assignment 2, PC-Astral Market, Y1, Y2, Y3'

def field(p,instruction):
    for tag,value in [('begin',None),('text',instruction),('separate',None),('end',None)]:
        r=p.add_run()
        if tag=='text':
            el=OxmlElement('w:instrText');el.set(qn('xml:space'),'preserve');el.text=' '+value+' '
        else:
            el=OxmlElement('w:fldChar');el.set(qn('w:fldCharType'),tag)
        r._r.append(el)

def para(text='',bold=False,size=None):
    p=doc.add_paragraph();p.paragraph_format.widow_control=True
    r=p.add_run(text);r.bold=bold
    if size:r.font.size=Pt(size)
    return p

def heading(text,level=2):return doc.add_heading(text,level)

def repository_link():
    url='https://github.com/ngminhtamtech-cmd/Assignment-WEB1044'
    p=doc.add_paragraph()
    link=OxmlElement('w:hyperlink')
    link.set(qn('r:id'),p.part.relate_to(url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');properties=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');properties.append(color)
    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');properties.append(underline)
    run.append(properties);text=OxmlElement('w:t');text.text=url;run.append(text)
    link.append(run);p._p.append(link)
    return p

def newpage(text,level=2):
    p=heading(text,level);p.paragraph_format.page_break_before=True
    p.paragraph_format.space_before=Pt(0)
    return p

figure_count=0
table_count=0

def figure(path,title,width=16):
    global figure_count
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Cm(width))
    p=doc.add_paragraph(style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('Hình ');field(p,'SEQ Hình \\* ARABIC');p.add_run('  '+title)
    figure_count+=1

def table(headers,rows,widths,title=None,font_size=11.5):
    global table_count
    if title:
        p=doc.add_paragraph(style='Caption');p.paragraph_format.keep_with_next=True
        p.add_run('Bảng ');field(p,'SEQ Bảng \\* ARABIC');p.add_run('  '+title);table_count+=1
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    table_borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'6');el.set(qn('w:color'),'D9D9D9');table_borders.append(el)
    t._tbl.tblPr.append(table_borders)
    for col,width in zip(t.columns,widths):col.width=Cm(width)
    for cell,text in zip(t.rows[0].cells,headers):cell.text=text
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for values in rows:
        for cell,text in zip(t.add_row().cells,values):cell.text=str(text)
    for ri,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'6');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            pr.append(borders);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF1' if ri==0 else 'FFFFFF');pr.append(shade)
            padding=OxmlElement('w:tcMar')
            for edge,value in [('top',85),('bottom',85),('left',100),('right',100)]:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(value));el.set(qn('w:type'),'dxa');padding.append(el)
            pr.append(padding)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                if ci==0 and widths[0]<2:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(font_size);r.bold=ri==0
    spacer=doc.add_paragraph()
    spacer.paragraph_format.line_spacing=1
    spacer.paragraph_format.space_after=Pt(0)
    spacer.paragraph_format.space_before=Pt(0)
    spacer.add_run(' ').font.size=Pt(5)
    return t

def screenshot(name,box):
    image=Image.open(ROOT/'docs/screenshots'/name)
    assert box[2]<=image.width and box[3]<=image.height,(name,image.size,box)
    path=QA/name;image.crop(box).save(path);return path

def centered(text,size=13,bold=False,before=0,after=6,style=None):
    p=doc.add_paragraph(style=style);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(before);p.paragraph_format.space_after=Pt(after);p.paragraph_format.line_spacing=1.35
    r=p.add_run(text);r.font.size=Pt(size);r.bold=bold;return p

# Cover ornament stays behind editable Word text and only on the cover.
p=sec.first_page_header.paragraphs[0]
inline=p.add_run().add_picture(str(frame_path),width=Cm(21),height=Cm(29.7))._inline
anchor=OxmlElement('wp:anchor')
for k,v in {'distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0','relativeHeight':'0','behindDoc':'1','locked':'0','layoutInCell':'1','allowOverlap':'1'}.items():anchor.set(k,v)
sp=OxmlElement('wp:simplePos');sp.set('x','0');sp.set('y','0');anchor.append(sp)
for kind in ['H','V']:
    pos=OxmlElement('wp:position'+kind);pos.set('relativeFrom','page')
    offset=OxmlElement('wp:posOffset');offset.text='0';pos.append(offset);anchor.append(pos)
for child in list(inline):
    if child.tag==qn('wp:extent'):anchor.append(child)
anchor.append(OxmlElement('wp:wrapNone'))
for child in list(inline):anchor.append(child)
inline.getparent().replace(inline,anchor)
centered('TRƯỜNG CAO ĐẲNG FPT POLYTECHNIC',13,True,before=22)
centered('MÔN WEB1044',12,True)
centered('LẬP TRÌNH CƠ SỞ VỚI JAVASCRIPT',12,True)
p=centered('',after=14)
for color in ['9869BF','C46DAB','7CA2B9','8CA673','E9B34B','E18D77','9CA7C3']:
    r=p.add_run('★ ');r.font.name='Arial';r.font.size=Pt(11);r.font.color.rgb=RGBColor.from_string(color)
p=centered('',before=8,after=0)
for letter,color in [('F','1466A3'),('P','F47D21'),('T','269344')]:
    r=p.add_run(letter);r.font.name='Arial';r.bold=True;r.italic=True;r.font.size=Pt(36);r.font.color.rgb=RGBColor.from_string(color)
r=p.add_run('  Education');r.font.name='Arial';r.font.size=Pt(10);r.font.color.rgb=RGBColor.from_string('666666')
p=centered('FPT POLYTECHNIC',14,True,after=32);p.runs[0].font.color.rgb=RGBColor.from_string('E66D19')
centered('BÁO CÁO ASSIGNMENT 2',19,True,after=15,style='Title')
centered('ĐỀ TÀI THIẾT KẾ WEBSITE BÁN MÁY TÍNH',13,True)
centered('PC ASTRAL MARKET',20,True,after=31,style='Title')
centered('Lớp: ................................................',13,True,after=18)
centered('Giảng viên hướng dẫn: Cô Lê Anh Đào',13,True,after=16)
centered('Nhóm sinh viên thực hiện',13,True,after=8)
table(['STT','Họ và tên','MSSV'],[['1','',''],['2','',''],['3','','']],[1.4,8.2,4.4],font_size=12)
centered('Cơ sở: ................................................',12,before=18,after=0)
centered('Năm 2026',12,True,before=38)
sec=doc.add_section(WD_SECTION_START.NEW_PAGE)
sec.left_margin=Cm(3);sec.right_margin=Cm(2);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2)
sec.different_first_page_header_footer=False
sec.header.is_linked_to_previous=False;sec.footer.is_linked_to_previous=False
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;field(p,'PAGE')
num=OxmlElement('w:pgNumType');num.set(qn('w:start'),'1');sec._sectPr.append(num)
doc.add_paragraph('Mục lục',style='Front Matter Title')
p=doc.add_paragraph();field(p,'TOC \\o "1-2" \\h \\z')
p=doc.add_paragraph('Danh mục hình',style='Front Matter Title');p.paragraph_format.page_break_before=True
p=doc.add_paragraph();field(p,'TOC \\h \\z \\c "Hình"')
doc.add_paragraph('Danh mục bảng',style='Front Matter Title')
p=doc.add_paragraph();field(p,'TOC \\h \\z \\c "Bảng"')

# Overview and requirements.
newpage('Chương 1 Giới thiệu đề tài',1)
heading('1.1 Mục tiêu của website')
para('Nhóm xây dựng PC-Astral Market để mô phỏng một website bán máy tính. Người xem có thể tìm sản phẩm, xem cấu hình, thêm vào giỏ và điền thông tin đặt hàng. Qua đề tài này, nhóm vận dụng JavaScript để xử lý dữ liệu, cập nhật giao diện, kiểm tra thông tin nhập và lưu giỏ trên trình duyệt.')
heading('1.2 Phạm vi thực hiện')
para('Website có 24 sản phẩm thuộc bốn danh mục: Laptop, PC để bàn, Màn hình và Phụ kiện. Mỗi danh mục có 6 sản phẩm. Ba nhóm mới, hot và khuyến mãi giúp giới thiệu sản phẩm; một sản phẩm có thể nằm trong nhiều nhóm.')
figure(FIGURES/'01-navigation.png','Sáu trang và luồng điều hướng')
para('Hình 1 thể hiện sáu trang của website. Người dùng có thể thêm hàng từ danh sách hoặc trang chi tiết, sau đó chuyển sang giỏ và đặt hàng. Website dùng dữ liệu tĩnh, không có máy chủ xử lý đơn, tài khoản hay thanh toán thật. Đơn hàng chỉ được lưu trong trình duyệt để minh họa chức năng.')
newpage('Chương 2 Phân tích yêu cầu',1)
para('Nhóm đối chiếu từng yêu cầu của đề bài với phần đã thực hiện trong Bảng 1. Các chức năng đều có thao tác thử trên giao diện, không chỉ trình bày hình ảnh.')
table(['Yêu cầu','Nội dung đã thực hiện'],[
('Y1.1\nKhởi tạo','Mở sáu trang, nạp dữ liệu và gắn các sự kiện. Không có lỗi JavaScript trong các luồng đã kiểm thử.'),
('Y1.2\nSản phẩm','24 sản phẩm từ mảng dữ liệu; có tên, giá, ảnh và nút thao tác. Có trang theo danh mục, ba nhóm mới/hot/khuyến mãi và hiệu ứng khi rê chuột.'),
('Y1.3\nSlider','Ba banner; có nút trước, tiếp, chấm chọn và tự chuyển bằng JavaScript.'),
('Y1.4\nKiểm tra form','Form đặt hàng kiểm tra thông tin rỗng, email và số điện thoại. Lỗi hiện cạnh ô cần sửa.'),
('Y1.5\nGiỏ hàng','Thêm, gộp sản phẩm trùng, sửa số lượng, xóa; lưu LocalStorage để tải lại trang vẫn còn giỏ.'),
('Y1.6\nNâng cao','Đồng hồ đếm ngược đến thời hạn khuyến mãi. Khi hết hạn, dùng lại giá thường.'),
('Y1.7\nTự đề xuất','Tìm kiếm, lọc và sắp xếp. Lý do lựa chọn và cách hoạt động trình bày ở mục 2.1.'),
('Y2\nTổ chức mã','HTML ở cấp gốc của website; JavaScript trong js; CSS trong css. Cấu trúc thực tế trình bày ở Chương 3.'),
('Y3\nBáo cáo','Có phân tích chức năng, sơ đồ, ảnh giao diện, kết quả kiểm thử và thông tin công cụ/thư viện.')
],[3.4,12.6],title='Đối chiếu yêu cầu Y1 đến Y3')
para('Countdown được chọn cho phần nâng cao. Ba chức năng tìm kiếm, lọc và sắp xếp được tính riêng cho phần tự đề xuất; không dùng lại giỏ hàng hoặc slider để thay cho yêu cầu này.')
newpage('2.1 Ba chức năng nhóm tự đề xuất')
heading('Tìm kiếm sản phẩm',3)
para('Nhóm chọn tìm kiếm vì người mua thường đã biết tên máy hoặc thương hiệu. Người dùng nhập từ khóa trên đầu trang; website bỏ khoảng trắng thừa và so sánh không phân biệt chữ hoa, chữ thường hay dấu tiếng Việt. Nếu không tìm thấy, trang hiện thông báo để người dùng đổi từ khóa.')
heading('Lọc sản phẩm',3)
para('Nhóm chọn bộ lọc để người mua thu hẹp lựa chọn theo danh mục, thương hiệu và ngân sách. Sản phẩm phải thỏa mãn đồng thời các điều kiện đã chọn. Khoảng giá gồm dưới 5 triệu, từ 5 đến dưới 15 triệu, từ 15 đến dưới 30 triệu và từ 30 triệu. Người dùng có thể xóa bộ lọc để xem lại toàn bộ danh sách.')
heading('Sắp xếp sản phẩm',3)
para('Nhóm chọn sắp xếp để người mua dễ so sánh giá. Có bốn cách: mặc định, giá tăng dần, giá giảm dần và tên A–Z. Chức năng chỉ đổi thứ tự của kết quả đang xem, không thay đổi dữ liệu sản phẩm gốc.')
figure(FIGURES/'02-selection.png','Thứ tự tìm kiếm lọc và sắp xếp')
para('Ví dụ: chọn Laptop, thương hiệu ASUS và mức giá từ 30 triệu sẽ còn hai sản phẩm ASUS TUF trong dữ liệu hiện tại. Chọn giá tăng dần sẽ đưa máy 35.990.000đ lên trước máy 38.990.000đ. Các lựa chọn nằm trên địa chỉ trang nên tải lại vẫn giữ được bộ lọc.')

# Current organization of source files.
newpage('Chương 3 Tổ chức và thiết kế',1)
heading('3.1 Tổ chức mã theo yêu cầu Y2')
para('Nhóm chia mã theo loại file và nhiệm vụ để dễ tìm chỗ sửa. Hình 3 thể hiện đúng thư mục website hiện tại. Sáu file HTML nằm cùng cấp với js, css và assets.')
figure(FIGURES/'03-folders.png','Cấu trúc thư mục website hiện tại')
para('Trong thư mục làm bài còn có docs chứa sơ đồ, ảnh minh chứng và kết quả kiểm thử; tools chứa công cụ hỗ trợ; tmp chứa file tạm; submission chứa bản nộp. Các thư mục hỗ trợ này nằm ngoài PC-Astral-Market. Gói nộp chỉ gồm thư mục website và báo cáo Word.')
newpage('3.2 Nhiệm vụ của các file')
para('HTML giữ bố cục, form và vùng hiển thị cố định. JavaScript đưa dữ liệu vào các vùng này và xử lý thao tác. Header, footer và thẻ sản phẩm được dùng chung để các trang có cách trình bày thống nhất.')
table(['File JavaScript','Nhiệm vụ chính'],[
('data.js','Lưu danh mục, sản phẩm, banner và thời hạn khuyến mãi.'),
('utils.js','Định dạng tiền, tìm sản phẩm, xác định giá, tìm kiếm, lọc, sắp xếp và kiểm tra thông tin người nhận.'),
('storage.js','Đọc và ghi dữ liệu trong trình duyệt; thông báo khi dữ liệu hỏng hoặc không lưu được.'),
('cart.js','Quản lý danh sách trong giỏ, số lượng và tổng tiền.'),
('common.js','Dựng phần giao diện dùng chung; cập nhật nút giỏ hàng, menu và thông báo.'),
('home.js, products.js,\nproduct-detail.js','Xử lý riêng cho trang chủ, danh sách và chi tiết sản phẩm.'),
('cart-page.js, checkout.js,\nabout.js','Xử lý riêng cho trang giỏ, đặt hàng và giới thiệu.')
],[6.3,9.7],title='Phân chia nhiệm vụ JavaScript')
heading('Cách nạp JavaScript',3)
para('Mỗi trang nạp theo thứ tự: dữ liệu → tiện ích → lưu trữ → giỏ hàng → giao diện chung → xử lý riêng của trang. Thuộc tính defer giúp các file chạy theo thứ tự sau khi HTML đã được đọc. Các file trao đổi qua một đối tượng chung tên PCAstral; không dùng import/export.')
heading('Cách chia CSS',3)
para('fonts.css khai báo font Sarabun; base.css đặt màu, chữ và quy tắc nền tảng; components.css định dạng thành phần dùng chung; pages.css định dạng bố cục từng trang; responsive.css điều chỉnh theo kích thước màn hình. Muốn sửa giao diện của một trang, nhóm tìm HTML tương ứng và phần CSS của trang đó trước.')
newpage('3.3 Dữ liệu của website')
para('Mỗi sản phẩm có một mã riêng để tra cứu. Hình 4 cho thấy mối liên hệ giữa danh mục, sản phẩm, dòng giỏ và đơn đã lưu. Đây là cách tổ chức dữ liệu JavaScript, không phải cơ sở dữ liệu SQL.')
figure(FIGURES/'04-data.png','Các nhóm dữ liệu và quan hệ')
para('Giỏ chỉ giữ mã sản phẩm và số lượng nên có thể lấy lại giá mới nhất. Đơn đã xác nhận giữ tên hàng và giá tại lúc đặt để không bị thay đổi theo giá hiển thị về sau.')
para('Danh mục giúp phân loại hàng hóa; mỗi dòng giỏ tham chiếu đến một sản phẩm. Khi xác nhận đơn, website chụp lại thông tin cần thiết để giữ đúng kết quả tại thời điểm đặt hàng. Cách tổ chức này giúp hạn chế việc lưu lặp tên, ảnh và giá trong giỏ.')

newpage('3.4 Use Case tổng quan')
para('Nhóm dùng sơ đồ Use Case để mô tả những việc khách hàng có thể thực hiện trên website. Khách hàng là người xem và mua hàng trong bản mô phỏng, không cần đăng nhập. Hình 5 thể hiện năm nhóm chức năng chính.')
figure(USECASES/'01-overview.png','Use Case tổng quan website')
para('Hình người bên ngoài khung là khách hàng; các hình elip bên trong là chức năng của website. Đường nối cho biết khách hàng được sử dụng chức năng đó. Hai sơ đồ tiếp theo trình bày cụ thể việc chọn sản phẩm, quản lý giỏ và đặt hàng.')
para('Xem xác nhận đơn gần nhất chỉ áp dụng cho đơn đã lưu trong phiên hiện tại. Website chưa có tài khoản, trang quản trị, lịch sử đơn hàng hoặc thanh toán trực tuyến.')

newpage('3.5 Use Case tìm và chọn sản phẩm')
para('Hình 6 trình bày sáu chức năng từ UC01 đến UC06. Mã UC giúp đối chiếu một chức năng giữa các sơ đồ; UC06 cũng xuất hiện ở phần giỏ hàng.')
figure(USECASES/'02-find-select.png','Use Case tìm và chọn sản phẩm')
para('Tìm kiếm (UC02), lọc (UC03) và sắp xếp (UC04) là ba chức năng nhóm tự đề xuất theo Y1.7. Người dùng có thể sử dụng riêng hoặc kết hợp; bộ lọc và sắp xếp vẫn hoạt động khi chưa nhập từ khóa.')
para('Người dùng có thể thêm giỏ ngay tại danh sách hoặc trang chi tiết. Khi không có kết quả, trang gợi ý đổi điều kiện tìm. Mã sản phẩm không tồn tại sẽ hiện thông báo; sản phẩm hết hàng không thể thêm giỏ.')

newpage('3.6 Use Case giỏ hàng và đặt hàng')
para('Hình 7 mô tả các thao tác với giỏ và đơn mô phỏng. UC06 giữ nguyên tên và mã từ sơ đồ trước. Khách hàng có thể xem giỏ, sửa số lượng hoặc xóa sản phẩm trước khi đặt.')
figure(USECASES/'03-cart-checkout.png','Use Case giỏ hàng và đặt hàng')
para('Mũi tên nét đứt có nhãn «include» chỉ bước kiểm tra luôn được dùng trong chức năng chính [7]. Thêm hàng và sửa số lượng đều kiểm tra giới hạn cho phép; đặt hàng luôn kiểm tra thông tin người nhận.')
para('Giỏ có hàng là điều kiện trước khi đặt đơn. Sau khi lưu thành công, người dùng xem được xác nhận gần nhất trong cùng phiên. Các trường hợp nhập sai hoặc không lưu được được mô tả tại mục 3.7.')

newpage('3.7 Mô tả các tình huống sử dụng chính')
def scenario(title,items):
    heading(title,3)
    for label,text in items:
        p=doc.add_paragraph();p.paragraph_format.line_spacing=1.15
        p.paragraph_format.space_after=Pt(4)
        p.add_run(label+': ').bold=True;p.add_run(text)

scenario('Tìm và chọn sản phẩm (UC01 đến UC06)',[
    ('Mục tiêu','Chọn được sản phẩm phù hợp nhu cầu và ngân sách.'),
    ('Điều kiện trước','Website đã hiển thị danh sách sản phẩm.'),
    ('Thao tác chính','Xem danh mục hoặc nhóm; tìm, lọc, sắp xếp; xem chi tiết hoặc thêm vào giỏ.'),
    ('Trường hợp cần xử lý','Không có kết quả thì đổi từ khóa hoặc xóa bộ lọc. Mã sai hiện thông báo; hết hàng thì không cho thêm.'),
    ('Kết quả','Hiển thị lựa chọn phù hợp; sản phẩm được thêm khi số lượng hợp lệ.'),
])
scenario('Quản lý giỏ hàng (UC06 đến UC09)',[
    ('Mục tiêu','Giữ đúng các sản phẩm và số lượng khách hàng muốn mua.'),
    ('Điều kiện trước','Có sản phẩm để thêm; thao tác sửa hoặc xóa cần có dòng trong giỏ.'),
    ('Thao tác chính','Thêm hàng, gộp mã trùng; sửa số lượng hoặc xóa; cập nhật tổng tiền và lưu giỏ.'),
    ('Trường hợp cần xử lý','Số lượng sai giữ giá trị cũ. Nếu bị chặn lưu, báo rõ giỏ chỉ được giữ tạm ở trang đang mở.'),
    ('Kết quả','Giỏ và tổng tiền đúng; tải lại vẫn giữ giỏ khi lưu thành công.'),
])
scenario('Đặt hàng mô phỏng (UC10 và UC11)',[
    ('Mục tiêu','Lưu đơn hợp lệ và xem lại xác nhận trong phiên hiện tại.'),
    ('Điều kiện trước','Giỏ có hàng; người mua đã kiểm tra sản phẩm và tổng tiền.'),
    ('Thao tác chính','Nhập thông tin, kiểm tra hợp lệ, lưu đơn; lưu thành công mới xóa giỏ và hiện xác nhận.'),
    ('Trường hợp cần xử lý','Giỏ rỗng chặn đặt. Form sai giữ nội dung và báo lỗi. Tổng tiền đổi cần xác nhận lại. Lưu lỗi giữ giỏ để thử lại; chặn gửi lặp.'),
    ('Kết quả','Có một đơn mô phỏng được lưu; tải lại vẫn xem được xác nhận trong cùng phiên.'),
])

newpage('3.8 Quản lý dự án qua GitHub và làm việc nhóm')
para('Nhóm lưu mã nguồn, tài nguyên, báo cáo và gói nộp tại kho GitHub chung. Lịch sử kho bắt đầu từ phiên bản dự án hiện tại; quy trình dưới đây áp dụng cho các lần cập nhật tiếp theo.')
repository_link()
heading('Phân công trong nhóm',3)
para('Nhóm có ba thành viên. Khi tiếp tục phát triển, nhóm chia việc theo ba phần: giao diện; xử lý JavaScript; kiểm thử và tổng hợp báo cáo. Trước khi sửa, các thành viên thống nhất phạm vi và người phụ trách. Mỗi phần cần có một bạn khác kiểm tra lại.')
heading('Quy trình phối hợp',3)
steps=[
    '1. Tạo Issue (đầu việc), nêu nội dung, người phụ trách và điều kiện hoàn thành. Có thể ghi kèm yêu cầu Y1, Y2 hoặc Y3 liên quan.',
    '2. Lấy phiên bản mới nhất từ nhánh main, sau đó tạo nhánh riêng cho phần việc đã nhận. Main giữ bản đã kiểm tra của nhóm.',
    '3. Sửa và thử chức năng liên quan; tạo commit với mô tả rõ việc đã làm, rồi đẩy nhánh lên GitHub.',
    '4. Mở Pull Request (đề nghị ghép thay đổi), ghi cách kiểm thử. Một thành viên khác xem lại trước khi ghép vào main.',
    '5. Sau khi ghép, các bạn cập nhật phiên bản mới. Người tổng hợp sửa báo cáo, ảnh minh chứng và gói ZIP nếu nội dung có thay đổi.',
]
for text in steps:
    p=para(text);p.paragraph_format.line_spacing=1.15
heading('Quản lý phiên bản và tài liệu',3)
para('Commit giúp theo dõi từng lần sửa và đối chiếu phiên bản khi có lỗi. Khi hai bạn sửa cùng một phần mã, nhóm trao đổi để xử lý xung đột và kiểm tra lại trước khi ghép. Không ghi đè lịch sử của nhánh chung.')
para('File Word do một bạn tổng hợp để tránh sửa đồng thời. Trước khi nộp, nhóm cập nhật mục lục, kiểm tra trang và tạo lại ZIP sao cho báo cáo bên trong khớp bản bên ngoài. README hướng dẫn chạy dự án; tệp tạm, cache và thông tin đăng nhập không đưa vào kho.')

# Interface evidence, with prose describing the customer's experience.
newpage('Chương 4 Hoạt động của website',1)
heading('4.1 Trang chủ và hiển thị sản phẩm')
figure(screenshot('01-home.png',(0,0,1440,1030)),'Trang chủ PC Astral Market')
para('Trang chủ có ô tìm kiếm, banner, bốn danh mục và ba nhóm sản phẩm mới, hot, khuyến mãi. Mỗi nhóm hiển thị bốn sản phẩm đại diện và có liên kết xem thêm. Thẻ sản phẩm lấy tên, ảnh, giá từ dữ liệu rồi đưa ra giao diện.')
para('Khi rê chuột hoặc di chuyển bằng bàn phím đến thẻ, thẻ được làm nổi bật và ảnh phóng nhẹ. Nút xem chi tiết và thêm giỏ luôn có thể thao tác trên màn hình cảm ứng.')
para('Giao diện sử dụng nền sáng, phần đầu trang màu đỏ và font Sarabun. Ảnh máy là ảnh thật được lưu trong dự án. Giá, nhóm hàng và tồn kho phục vụ bài thực hành; không phải thông tin bán hàng được cập nhật trực tiếp.')
newpage('4.2 Danh sách và chi tiết sản phẩm')
figure(screenshot('02-products-filtered.png',(100,400,1340,968)),'Danh sách sau khi lọc Laptop ASUS')
para('Trang danh sách hiển thị số kết quả cùng các điều kiện đang chọn. Tìm kiếm và bộ lọc kết hợp với nhau như Hình 2; người dùng có thể tiếp tục sắp xếp mà không mất lựa chọn.')
figure(screenshot('03-product-detail.png',(100,190,1340,840)),'Thông tin chi tiết một sản phẩm')
para('Trang chi tiết lấy đúng sản phẩm theo mã trên địa chỉ trang. Nếu mã sai hoặc thiếu, trang báo không tìm thấy. Nếu hết hàng, nút mua bị khóa. Số lượng thêm vào giỏ phải nằm trong tồn kho minh họa.')
newpage('4.3 Quản lý trạng thái và tính tổng')
para('Trạng thái giỏ hàng là danh sách sản phẩm người dùng đã chọn và số lượng tương ứng. Khi mở trang, website đọc lại giỏ đã lưu. Khi người dùng thêm, sửa hoặc xóa, website cập nhật danh sách, tính lại tiền rồi hiển thị kết quả theo Hình 11.')
figure(FIGURES/'06-cart-state.png','Luồng quản lý giỏ và tính tổng tiền')
table(['Ví dụ trong thời gian khuyến mãi','Kết quả'],[
('ASUS TUF Gaming 16, đơn giá 38.990.000đ; mua 3 máy','38.990.000 × 3 = 116.970.000đ'),
('Phí giao hàng trong bản mô phỏng','0đ'),
('Tổng thanh toán; số lượng trên nút giỏ hàng','116.970.000đ; 3 sản phẩm')
],[9,7],title='Ví dụ tính tổng giỏ hàng')
newpage('4.4 Lưu giỏ và xử lý dữ liệu không hợp lệ')
figure(screenshot('04-cart.png',(100,190,1340,850)),'Giỏ hàng sau khi đổi số lượng và tải lại')
para('Thêm cùng một sản phẩm sẽ tăng số lượng trên dòng đã có. Khi sửa số lượng, website chỉ nhận số nguyên từ 1 đến tồn kho. Nếu nhập 0, số âm, số lẻ hoặc quá tồn kho, trang hiện lỗi và trả về số lượng hợp lệ trước đó. Muốn bỏ sản phẩm, người dùng bấm Xóa.')
para('Giỏ được lưu bằng LocalStorage. Khi chuyển trang hoặc tải lại trên cùng địa chỉ và cổng chạy website, giỏ vẫn được giữ. Nếu giỏ rỗng, trang hiện lời nhắc quay lại mua sắm và không cho đặt hàng.')
heading('Khi dữ liệu lưu có vấn đề',3)
para('Website kiểm tra dữ liệu trước khi dùng: bỏ mã sản phẩm không còn tồn tại, bỏ số lượng sai, gộp dòng trùng và giới hạn theo tồn kho. Nếu dữ liệu lưu bị hỏng, website thông báo và dùng giỏ mặc định thay vì dừng hoạt động.')
para('Nếu trình duyệt không cho lưu, website báo rõ giỏ chỉ được giữ tạm tại trang đang mở. Thay đổi giỏ từ tab khác cũng được cập nhật. Thông tin tên và giá không lưu lặp trong giỏ mà được lấy từ danh sách sản phẩm.')

# Explanatory logic images replace all source-code excerpts.
newpage('4.5 Cách hoạt động của slider')
para('Slider có ba banner. Website nhớ banner đang hiển thị; mỗi thao tác của người dùng hoặc lượt tự chuyển sẽ chọn banner cần hiển thị tiếp theo. Hình 13 thể hiện hai cách điều khiển này.')
figure(FIGURES/'07-slider.png','Luồng điều khiển slider')
para('Nút Trước và Tiếp chuyển lần lượt giữa ba banner. Chấm chọn đưa người dùng đến đúng banner tương ứng. Khi chuyển, cả ảnh, tiêu đề, mô tả và liên kết đều thay đổi cùng nhau.')
para('Cứ 6 giây, website kiểm tra có thể tự chuyển hay không. Nếu chuột đang ở trên slider, bàn phím đang chọn phần tử trong slider, người dùng bấm Tạm dừng hoặc tab đang bị ẩn, banner được giữ nguyên. Thao tác này giúp người xem có thời gian đọc và bấm liên kết.')
newpage('4.6 Cách hoạt động của countdown')
para('Đồng hồ đếm ngược cho biết thời gian còn lại của chương trình khuyến mãi. Nhóm đặt thời hạn cố định trong dữ liệu là 23 giờ 59 phút 59 giây ngày 30/10/2026 theo giờ Việt Nam.')
figure(FIGURES/'08-countdown.png','Đếm ngược và xử lý khi hết khuyến mãi')
para('Mỗi giây, website lấy thời hạn trừ thời gian hiện tại rồi đổi kết quả thành ngày, giờ, phút và giây. Khi tải lại trang hoặc quay lại tab, website tính từ mốc kết thúc cũ nên thời gian không bị bắt đầu lại.')
para('Khi không còn thời gian, đồng hồ dừng và hiện “Chương trình đã kết thúc”. Giá ưu đãi hết hiệu lực; danh sách, chi tiết, giỏ và phần đặt hàng cùng dùng giá thường. Giá được xác định theo cùng một cách ở mọi trang để tránh lệch tổng tiền.')
para('Trước khi lưu đơn, website kiểm tra lại tổng tiền. Nếu tổng vừa thay đổi, người dùng được yêu cầu xem lại và xác nhận lần nữa. Các thông tin đang nhập trong form vẫn được giữ.')
newpage('4.7 Luồng đặt hàng và kiểm tra thông tin')
para('Form nhận họ tên, email, số điện thoại, địa chỉ và ghi chú tùy chọn. Bốn trường đầu bắt buộc nhập; khoảng trắng không được tính là thông tin hợp lệ. Hình 15 mô tả các bước kiểm tra trước khi xác nhận.')
figure(FIGURES/'09-checkout.png','Luồng đặt đơn và các trường hợp cần sửa')
para('Email phải đúng định dạng cơ bản. Điện thoại nhận 10 chữ số bắt đầu bằng 0 hoặc dạng +84 tương ứng. Lỗi hiện cạnh từng ô và con trỏ chuyển đến ô lỗi đầu tiên. Website chặn việc gửi lặp trong lúc đang xử lý.')
newpage('4.8 Minh chứng kiểm tra form và xác nhận đơn')
figure(screenshot('05-checkout-errors.png',(100,190,1340,990)),'Form thông báo thông tin chưa hợp lệ',width=14.2)
para('Ở Hình 16, người dùng cần sửa các ô được báo lỗi trước khi tiếp tục. Chỉ phần tổng tiền được cập nhật khi giỏ hoặc giá thay đổi, vì vậy nội dung đang nhập không bị xóa.')
figure(screenshot('07-order-confirmation.png',(100,190,1340,880)),'Xác nhận đơn hàng mô phỏng',width=14.2)
para('Website lưu đơn vào SessionStorage trước, sau đó mới xóa giỏ. Nếu lưu thất bại, giỏ được giữ để thử lại. Khi lưu thành công, tải lại trang xác nhận trong cùng phiên vẫn xem được đơn gần nhất; không có giao dịch hay email thật.')
newpage('4.9 Trang giới thiệu và giao diện trên điện thoại')
figure(screenshot('08-about.png',(100,190,1340,850)),'Trang giới thiệu và chính sách minh họa')
para('Trang giới thiệu trình bày mục tiêu cửa hàng, phạm vi bài tập và chính sách minh họa. Các thông tin này giúp người xem hiểu đây là website mô phỏng, không phải một cửa hàng đang nhận đơn thật.')
mobile=Image.new('RGB',(850,780),'white')
for x,name in [(0,'09-mobile-home.png'),(460,'10-mobile-products.png')]:
    im=Image.open(ROOT/'docs/screenshots'/name)
    if name=='09-mobile-home.png':
        im=im.crop((0,0,390,748))
    else:
        im=im.crop((0,430,390,1025))
    mobile.paste(im,(x,0))
mobile_path=QA/'mobile-pair.png';mobile.save(mobile_path)
figure(mobile_path,'Trang chủ và bộ lọc trên màn hình điện thoại',width=10)
para('Trên màn hình nhỏ, menu và bộ lọc có nút đóng/mở; các khối được sắp lại theo chiều dọc. Website đã được kiểm tra ở chiều rộng 390, 768 và 1440 px, đồng thời kiểm tra thao tác bằng phím Tab và Enter.')

# Recorded checks only; this report edit does not claim a new full test run.
record=json.loads((ROOT/'docs/test-results.json').read_text(encoding='utf-8-sig'))
assert len(record['results'])==18 and all(r['status']=='PASS' for r in record['results'])
extra=json.loads((ROOT/'docs/refactor-tests.json').read_text(encoding='utf-8-sig'))
assert len(extra['results'])==7 and all(r['status']=='PASS' for r in extra['results'])
newpage('Chương 5 Kiểm thử website',1)
para('Kết quả dưới đây được ghi nhận ngày 17/10/2026 trên Chrome và Edge, với website chạy bằng HTTP cục bộ. Các thao tác được chạy bằng Playwright; công cụ này chỉ phục vụ kiểm thử, không được nạp vào website.')
para('Bộ kiểm thử chính có 18 nhóm và đều đạt. Bảng 4 và Bảng 5 tóm tắt thao tác cùng kết quả quan sát được. Nhật ký chi tiết nằm trong docs/test-results.json và docs/refactor-tests.json.')
table(['STT','Thao tác kiểm tra','Kết quả thực tế'],[
('01','Mở trang chủ và xem các nhóm sản phẩm','Đủ 24 sản phẩm, 4 danh mục và 3 nhóm. Đạt.'),
('02','Bấm trước, tiếp, chấm chọn; chờ tự chuyển; chọn phần tử trong slider','Chuyển đúng banner; tự chuyển và tạm dừng đúng. Đạt.'),
('03','Tìm từ khóa khác chữ hoa, dấu và khoảng trắng','Tìm đúng sản phẩm; dữ liệu gốc không đổi. Đạt.'),
('04','Kết hợp bộ lọc, sắp xếp rồi tải lại; thử giá ở ranh giới','Kết quả đúng điều kiện; giữ lựa chọn; xóa lọc về 24 sản phẩm. Đạt.'),
('05','Mở chi tiết với mã đúng, mã sai, thiếu mã; thử hết hàng và lỗi ảnh','Hiển thị đúng; báo không tìm thấy; khóa mua khi hết hàng; có chữ thay ảnh lỗi. Đạt.'),
('06','Thêm trùng; sửa số lượng; xóa; tải lại giỏ','Gộp dòng đúng, tổng đúng, giữ giỏ sau tải lại; từ chối số lượng sai. Đạt.'),
('07','Đưa vào giỏ dữ liệu lưu hỏng, mã lạ và dòng trùng','Không dừng trang; bỏ dữ liệu sai; gộp và giới hạn số lượng. Đạt.'),
('08','Thay đổi giỏ ở tab thứ hai','Giỏ ở tab thứ nhất được cập nhật. Đạt.'),
('09','Gửi form rỗng, toàn dấu cách, email và điện thoại sai','Hiện lỗi ở từng ô; chuyển con trỏ đến ô lỗi đầu. Đạt.')
],[1.2,6.5,8.3],title='Kiểm thử sản phẩm slider và giỏ hàng',font_size=11)
newpage('5.1 Kiểm thử đặt hàng và các tình huống bổ sung')
table(['STT','Thao tác kiểm tra','Kết quả thực tế'],[
('10','Nhập hợp lệ, thử số +84, gửi lặp và tải lại xác nhận','Lưu một đơn, xóa giỏ sau khi lưu, giữ xác nhận. Đạt.'),
('11','Chặn lưu SessionStorage khi đặt hàng','Báo chưa lưu được; giữ giỏ và cho thử lại. Đạt.'),
('12','Chặn lưu LocalStorage','Có cảnh báo; giỏ chỉ được giữ tạm tại trang đang mở. Đạt.'),
('13','Mô phỏng thời gian đi qua hạn khuyến mãi','Dừng đồng hồ; dùng giá thường và tính lại giỏ. Đạt.'),
('14','Xem sáu trang ở 390, 768 và 1440 px','Không tràn ngang; menu và bộ lọc nhỏ hoạt động. Đạt.'),
('15','Dùng Tab và Enter để điều hướng, tìm kiếm','Thao tác được bằng bàn phím. Đạt.'),
('16','Theo dõi lỗi JavaScript và tài nguyên tải','Không có lỗi trong các luồng kiểm thử. Đạt.'),
('17','Thử mua hàng, tải lại và form trên Edge','Giỏ được giữ; form và bố cục hoạt động đúng. Đạt.'),
('18','Mở 24 ảnh sản phẩm và kiểm tra nguồn','Ảnh tải từ dự án; đủ thông tin nguồn; website không dùng SVG. Đạt.')
],[1.2,6.5,8.3],title='Kiểm thử đặt hàng lưu trữ và giao diện',font_size=11)
heading('5.2 Kiểm tra sau khi sắp xếp lại mã')
para('Bảy nhóm kiểm tra bổ sung đều đạt. Các điểm được kiểm tra lại gồm: thứ tự nạp file, giữ con trỏ khi sửa giỏ, giữ nội dung form khi giá thay đổi, lưu đơn trước khi xóa giỏ, chặn gửi lặp và kiểm tra sáu trang trên hai trình duyệt.')
para('Khi kiểm tra hạn khuyến mãi, công cụ mô phỏng thời gian để quan sát cả trước và sau thời hạn. Không cần chờ đến ngày kết thúc hoặc thay đổi đồng hồ của máy tính.')

newpage('Chương 6 Kết luận',1)
para('Nhóm đã hoàn thành website PC-Astral Market với sáu trang và luồng mua hàng mô phỏng từ chọn sản phẩm đến xác nhận đơn. Các yêu cầu Y1, Y2 và Y3 được thể hiện qua chức năng hoạt động, cách chia thư mục và nội dung báo cáo.')
heading('6.1 Giới hạn và hướng phát triển')
para('Dữ liệu sản phẩm là dữ liệu tĩnh; giá không tự cập nhật từ cửa hàng. Tồn kho và khuyến mãi là minh họa. Giỏ và đơn lưu trên trình duyệt nên có thể bị xóa, không dùng để quản lý bán hàng thực tế.')
para('Nếu phát triển tiếp, nhóm có thể bổ sung máy chủ lưu đơn, quản lý sản phẩm và kiểm tra tồn kho. Đăng nhập, quản trị và thanh toán trực tuyến chưa thuộc phạm vi bài này.')
heading('6.2 Công nghệ và công cụ sử dụng')
table(['Nhóm','Công nghệ hoặc công cụ','Vai trò'],[
('Website','HTML5, CSS3, JavaScript thuần','Bố cục, giao diện và xử lý chức năng.'),
('Trình duyệt','DOM, LocalStorage, SessionStorage, bộ hẹn giờ, URLSearchParams','Cập nhật nội dung, lưu dữ liệu, điều khiển thời gian và đọc lựa chọn.'),
('Thư viện website','Không sử dụng thư viện bên ngoài','Không cài framework, thư viện slider hay kiểm tra form.'),
('Chạy và kiểm thử','HTTP cục bộ, Chrome, Edge, Playwright','Mở website, kiểm tra thao tác và chụp minh chứng. Có thể chạy bằng VS Code Live Server.'),
('Tài liệu','Microsoft Word, python-docx, Pillow','Soạn báo cáo, vẽ sơ đồ và kiểm tra trình bày.')
],[2.7,6.1,7.2],title='Công nghệ API và công cụ sử dụng',font_size=11)
para('LocalStorage và SessionStorage là khả năng có sẵn của trình duyệt, không phải thư viện cần cài. Font Sarabun và ảnh sản phẩm được lưu trong dự án nên khi demo không phải tải từ CDN.')
newpage('Phụ lục Hướng dẫn chạy và tài liệu tham khảo',1)
heading('Hướng dẫn chạy website',3)
para('Giải nén gói nộp, mở thư mục PC-Astral-Market bằng VS Code, chọn index.html rồi chọn Open with Live Server. Chuyển trang bằng các liên kết trong website và giữ cùng địa chỉ, cổng chạy để dùng chung giỏ đã lưu.')
para('Nếu trình duyệt hiện một website khác, dừng Live Server đang chạy và mở lại đúng thư mục PC-Astral-Market. Không mở các file HTML bằng đường dẫn lấy từ một dự án khác.')
heading('Trình tự trình bày khi bảo vệ',3)
for text in ['1. Mở trang chủ, giới thiệu ba nhóm sản phẩm, điều khiển slider và xem đồng hồ đếm ngược.','2. Tìm kiếm, lọc và sắp xếp; mở chi tiết rồi thêm cùng sản phẩm hai lần để chứng minh gộp số lượng.','3. Sửa số lượng, xóa một dòng, tải lại trang để kiểm tra lưu giỏ.','4. Gửi form sai để xem lỗi; nhập hợp lệ rồi xác nhận đơn mô phỏng và kiểm tra giỏ đã được xóa.']:
    p=para(text);p.paragraph_format.line_spacing=1.15
para('Thời hạn khuyến mãi hiện tại là ngày 30/10/2026. Nếu demo sau thời điểm này, website sẽ hiển thị chương trình kết thúc. Khi cần minh họa lại ưu đãi, đặt một thời hạn tương lai cố định trong data.js trước buổi trình bày.')
heading('Tài liệu tham khảo và nguồn ảnh',3)
refs=['[1] Đề Assignment Lập trình cơ sở với JavaScript, môn WEB1044, tài liệu môn học gồm 5 trang.','[2] PC Market, https://pcmarket.vn/. Nguồn tên máy, cấu hình, ảnh và giá tham khảo thu thập ngày 16/10/2026. Liên kết từng sản phẩm có trong ASSET-SOURCES.md.','[3] MDN Web Docs, phần tử script và thuộc tính defer: https://developer.mozilla.org/en-US/docs/Web/HTML/Element/script','[4] MDN Web Docs, lưu trữ trong trình duyệt: https://developer.mozilla.org/en-US/docs/Web/API/Web_Storage_API','[5] MDN Web Docs, kiểm tra dữ liệu form: https://developer.mozilla.org/en-US/docs/Learn_web_development/Extensions/Forms/Form_validation','[6] Font Sarabun, giấy phép SIL Open Font License lưu tại assets/fonts/sarabun/OFL.txt.']
refs.append('[7] IBM Documentation, quan hệ include trong sơ đồ Use Case: https://www.ibm.com/docs/en/dma?topic=diagrams-include-relationships')
for text in refs:
    p=para(text,size=10.5);p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(5)
para('Ảnh sản phẩm thuộc nguồn được ghi nhận; ảnh giao diện trong báo cáo được chụp từ website của đề tài.',size=11)
for style in doc.styles:
    if style.type in (1,2):
        fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in list(fonts.attrib):
            if 'theme' in attr.lower():del fonts.attrib[attr]
        for attr in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+attr),'Times New Roman')
    for border in style.element.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)
report=OUT/'BaoCao_Assignment2_PC-Astral-Market.docx';doc.save(report)
manifest={'report':str(report),'figures':figure_count,'tables':table_count,'usecase_diagrams':3,'toc_levels':2,'toc_leader':'dots','personal_fields':'3 member names and student IDs left blank','instructor':'Lê Anh Đào','code_excerpts':0,'structure_checked_against_files':True}
(ROOT/'docs/report-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(manifest,ensure_ascii=True,indent=2))
